"""Rebuild submission workbooks from measured evidence and public-view templates.

The template owner disabled original downloads. This script reconstructs the
visible sheet names and column headers; it does not claim to modify the originals.
All measured results retain their true status. No peer review/signature is invented.
"""
from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "submission"
PRIVATE = SUB / "private"
SELF = PRIVATE / "self-test"
EVIDENCE = SUB / "evidence"
SCOPE = "Tự kiểm thử hệ thống của mình; chưa thay thế phần tester bạn cặp."
SOURCE = "Tái dựng theo tên sheet/cột từ chế độ xem công khai của mẫu Drive; file gốc bị khóa tải xuống."
DRIVE = "https://drive.google.com/drive/folders/173jy7ytRCeVvkClLM5JoPLcjbKE2jT9Q"
SITE = "https://sv23.bcse-vju.com"
REPO = "https://github.com/minhhal1180/qlychtsv"
IDENTITY = "Bùi Xuân Minh"
MSSV = "23110193"
LOCAL_TZ = timezone(timedelta(hours=7))
NAVY, GREEN, LIGHT, YELLOW, RED = "173A53", "176D53", "EAF2EE", "FFF6D5", "FCE9E8"


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "\n".join(f"{i}. {text(item)}" for i, item in enumerate(value, 1))
    return json.dumps(value, ensure_ascii=False, indent=2)


def date_of(value):
    if not value:
        return "Chưa xác định"
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(LOCAL_TZ).strftime("%d/%m/%Y")


def partner_info():
    path = SUB / "partner.json"
    if not path.exists():
        return "Chưa được cung cấp; chưa có kiểm thử chéo"
    data = read(path)
    return text(data)


def relative_evidence(value):
    value = str(value).replace("\\", "/")
    if value.startswith("submission/"):
        value = value[len("submission/"):]
    return value


def workbook_link(value, self_test=False):
    prefix = "../" if self_test else ""
    return prefix + relative_evidence(value)


FUNCTIONS = {
    "FR-01":"Đăng nhập", "FR-02":"Phiên và đăng xuất", "FR-03":"Quản lý tài khoản",
    "FR-04":"Vai trò và khóa", "FR-05":"Danh mục loại", "FR-06":"Tạo yêu cầu",
    "FR-07":"Danh sách và tìm kiếm", "FR-08":"Chi tiết và quyền sở hữu", "FR-09":"Hủy trước tiếp nhận",
    "FR-10":"Tiếp nhận", "FR-11":"Yêu cầu bổ sung", "FR-12":"Sinh viên bổ sung",
    "FR-13":"Trao đổi và tệp", "FR-14":"Hoàn tất và từ chối", "FR-15":"Mở lại", "FR-16":"Thông báo nội bộ",
}
NFR_GROUPS = {"NFR-01":"Hiệu năng", "NFR-02":"Hiệu năng", "NFR-03":"Hiệu năng", "NFR-04":"Hiển thị",
              "NFR-05":"Nội dung", "NFR-06":"Bảo mật", "NFR-07":"Nhất quán / retry",
              "NFR-08":"Chịu dữ liệu lỗi", "NFR-09":"Giới hạn tệp", "NFR-10":"Truy vết"}
E2E_SRS_FLOWS = {
    "TC-FLOW01-01":"FLOW-01", "TC-FLOW01-02":"FLOW-01", "TC-FLOW01-03":"FLOW-02",
    "TC-FLOW02-01":"FLOW-01 / nhánh hủy", "TC-FLOW02-02":"FLOW-03 / nhánh từ chối",
    "TC-FLOW02-03":"FLOW-03 / nhánh từ chối", "TC-FLOW03-01":"FLOW-03 / mở lại",
    "TC-FLOW03-02":"FLOW-03 / cán bộ cũ không hợp lệ", "TC-FLOW03-03":"FLOW-03 + FLOW-02",
}
TECHNIQUES = ["Phân hoạch tương đương (EP)", "Giá trị biên (BVA)", "Bảng quyết định", "Chuyển trạng thái",
              "Use case / E2E", "Phân quyền", "Error guessing / Checklist", "Exploratory"]
SEVERITIES = ["Critical", "Major", "Minor", "Trivial"]
CATEGORIES = ["Functional", "Visual", "Content", "Performance", "Crash"]
STATUSES = ["Pass", "Fail", "Blocked", "Chưa chạy"]
BUG_STATES = ["Mới", "Đã xác nhận", "Bị từ chối", "Trùng lặp", "Cần thêm thông tin"]
TEMPLATE_HEADERS = {
    "Review_SRS":["STT","Vị trí trong SRS (mục / mã yêu cầu)","Nội dung trích từ SRS","Vấn đề phát hiện","Loại vấn đề","Mức độ","Đề xuất sửa","Phản hồi của chủ hệ thống","Lý do / cách đã sửa"],
    "TC_Chuc_nang":["TC ID","Chức năng / Luồng","Kỹ thuật thiết kế","Tiêu đề","Tiền điều kiện","Dữ liệu thử","Các bước thực hiện","Kết quả mong đợi","Độ ưu tiên","Kết quả","Kết quả thực tế (khi Fail)","Bug ID","Ghi chú"],
    "TC_Phi_chuc_nang":["TC ID","Mã NFR","Nhóm","Tiêu đề","Công cụ / phương pháp","Điều kiện / môi trường","Các bước thực hiện","Ngưỡng đạt (theo SRS)","Giá trị đo được","Kết quả","Bug ID","Ghi chú"],
    "Test_Log":["STT","Ngày","Giờ (bắt đầu – kết thúc)","Link / phiên bản","Trình duyệt / thiết bị","Phạm vi đã chạy (TC ID)","Số TC chạy","Pass","Fail","Blocked","Bug mới (ID)","Sự cố / ghi chú"],
    "Bug_Report":["Bug ID","Tiêu đề","Loại lỗi","Mức nghiêm trọng","Mức ưu tiên","Chức năng / NFR","TC liên quan","URL / màn hình","Trình duyệt / thiết bị","Tiền điều kiện","Các bước tái hiện","Kết quả thực tế","Kết quả mong đợi (trích SRS)","Bằng chứng (link)","Ngày phát hiện","Trạng thái","Phản hồi của chủ hệ thống"],
    "Bug_Manifest":["Mã lỗi","Loại lỗi","Mức nghiêm trọng","Chức năng / Luồng","Vị trí (màn hình, URL, API)","Yêu cầu SRS bị vi phạm (mã + trích nội dung)","Mô tả lỗi","Các bước tái hiện","Kết quả mong đợi (theo SRS)","Kết quả thực tế","Bằng chứng (link ảnh / video lỗi trên sản phẩm)","File chứa lỗi","Ghi chú"],
}


def technique(row):
    value = row.get("technique", "")
    return {"Equivalence partitioning":"Phân hoạch tương đương (EP)", "Boundary value analysis":"Giá trị biên (BVA)",
            "Decision table / quyền":"Bảng quyết định", "State transition":"Chuyển trạng thái"}.get(value, value or "Error guessing / Checklist")


def result_status(value):
    return {"PASS":"Pass", "FAIL":"Fail", "ERROR":"Blocked", "BLOCKED":"Blocked", "NOT_RUN":"Chưa chạy"}.get(value, value)


def build_catalog(functional, nfr):
    core = read(EVIDENCE / "seeded_functional" / "results.json")
    extra = read(EVIDENCE / "additional" / "results.json")
    exploratory = read(EVIDENCE / "exploratory" / "results.json")
    fr_map = {"B01":"FR-06", "B02":"FR-05", "B03":"FR-09", "B04":"FR-14", "B05":"FR-14",
              "B06":"FR-12", "B07":"FR-15", "B08":"FR-07", "B09":"FR-16", "B10":"FR-14 / NFR-10"}
    requirement_quotes = {
        "B01":"FR-06: Tiêu đề1–200 sau trim; toàn khoảng trắng phải400, không tạo hồ sơ.",
        "B02":"FR-05: Chỉ dùng loại active, kể cả khi loại bị tắt sau khi mở form.",
        "B03":"FR-09: Chỉ NEW được hủy; đã tiếp nhận phải409 và giữ IN_PROGRESS.",
        "B04":"FR-14: COMPLETE bắt buộc kết quả1–10000 ký tự sau trim.",
        "B05":"FR-14: REJECT bắt buộc lý do1–10000 ký tự sau trim.",
        "B06":"FR-12: Bổ sung hợp lệ từ WAITING_STUDENT phải chuyển về IN_PROGRESS.",
        "B07":"FR-15: Chỉ COMPLETED mở lại; REJECTED phải409.",
        "B08":"FR-07: Tìm chuỗi con theo mã HOẶC tiêu đề.",
        "B09":"FR-16: Sau đánh dấu đã đọc và refresh, is_read vẫn1.",
        "B10":"FR-14 / NFR-10: Mỗi chuyển trạng thái có actor,time,from,to; có sự kiện COMPLETE.",
    }
    file_map = {"B01":"app.py:create_request", "B02":"app.py:create_request", "B03":"app.py:action(cancel)",
                "B04":"app.py:action(complete)", "B05":"app.py:action(reject)", "B06":"app.py:message",
                "B07":"app.py:action(reopen)", "B08":"app.py:list_requests", "B09":"app.py:read_notification", "B10":"app.py:transition"}
    url_map = {"B01":"POST /api/requests", "B02":"POST /api/requests", "B03":"POST /api/requests/:id/action (cancel)",
               "B04":"POST /api/requests/:id/action (complete)", "B05":"POST /api/requests/:id/action (reject)",
               "B06":"POST /api/requests/:id/messages", "B07":"POST /api/requests/:id/action (reopen)",
               "B08":"GET /api/requests?q=HT-00006", "B09":"POST /api/notifications/:id/read", "B10":"Chi tiết yêu cầu / lịch sử"}
    bugs=[]
    for row in sorted(core["results"], key=lambda r:r["id"]):
        bug_id=row["id"]
        major=bug_id in {"B02","B03","B04","B05","B06","B07","B10"}
        tests=[r["id"] for r in functional if r.get("bug_id")==bug_id]
        tests.extend(r["id"] for r in nfr if r.get("bug_id")==bug_id)
        evidence=[f"evidence/seeded_functional/{row['screenshot']}",f"evidence/seeded_functional/{bug_id}.json"]
        bugs.append({"id":bug_id,"title":row["title"],"category":"Functional","severity":"Major" if major else "Minor",
                     "priority":"High" if major else "Medium","requirement":fr_map[bug_id],"requirement_quote":requirement_quotes[bug_id],
                     "spec_relation":"Liên quan đặc tả SRS v1.1; truy vết F01–F05/PDF gốc theo mục tương ứng.","seeded":True,
                     "url":url_map[bug_id],"preconditions":"Tài khoản mẫu active; database demo độc lập cho từng phép thử; TRAINING_BUGS=1.",
                     "steps":row["steps"],"expected":row["expected"],"actual":row["actual"],"testcases":tests,
                     "evidence":evidence,"file":file_map[bug_id],"date":date_of(core["finished_at_utc"]),
                     "environment":"Windows / Edge155, viewport1440×1100; HTTP và ảnh UI trên database tạm.",
                     "status":"Đã xác nhận","owner_response":"Đã xác nhận nội bộ bằng assertion và ảnh; giữ lỗi cho demo, chưa có phản hồi bạn cặp."})
    for row in sorted((r for r in extra["results"] if r["mode"]=="training"),key=lambda r:r["id"]):
        bug_id=row["id"]
        major=row["category"] in {"Performance","Crash"} or bug_id=="C01"
        evidence=[f"evidence/additional/{name}" for name in row.get("evidence_files",[])]
        json_path=f"evidence/additional/{bug_id}-training.json"
        if json_path not in evidence:evidence.append(json_path)
        # Screenshot reports may be added later without changing the generator.
        for candidate in [f"{bug_id}-training.png",f"{bug_id}-report.png"]:
            if (EVIDENCE/"additional"/candidate).exists():
                candidate_path=f"evidence/additional/{candidate}"
                if candidate_path not in evidence:evidence.insert(0,candidate_path)
        endpoint=row.get("actual",{}).get("endpoint") if isinstance(row.get("actual"),dict) else None
        location=endpoint or {"V01":"Chi tiết yêu cầu / tiêu đề", "V02":"Hộp thoại Tạo yêu cầu,375px", "V03":"Danh sách / nhãn trạng thái,375px",
                              "C01":"Danh sách / nhãn COMPLETED", "C02":"Tạo yêu cầu / hướng dẫn tệp", "C03":"Danh sách / Ngày tạo"}.get(bug_id,"Giao diện/API")
        file_name="static/training.css" if row["category"]=="Visual" else "static/training.js" if row["category"]=="Content" else "training_extras.py"
        bugs.append({"id":bug_id,"title":row["title"],"category":row["category"],"severity":"Major" if major else "Minor",
                     "priority":"High" if major else "Medium","requirement":row["requirement"],"requirement_quote":f"{row['requirement']}: {row['expected']}",
                     "spec_relation":"Yêu cầu bổ sung trong SRS v1.1 cho bài thực hành; không gán là yêu cầu đã phê duyệt trong PDF gốc.","seeded":True,
                     "url":location,"preconditions":"Tài khoản mẫu active; TRAINING_BUGS=1; dữ liệu mẫu5yêu cầu; dùng viewport/số mẫu theo các bước.",
                     "steps":row["steps"],"expected":row["expected"],"actual":row["actual"],
                     "testcases":[r["id"] for r in nfr if r.get("bug_id")==bug_id],"evidence":evidence,"file":file_name,
                     "date":date_of(extra["finished_at_utc"]),"environment":"Windows / Edge155.0.4283.62 / HTTP cục bộ; không phải đo tải VPS.",
                     "status":"Đã xác nhận","owner_response":"Đã xác nhận nội bộ bằng so sánh baseline/training; giữ lỗi cho demo, chưa có phản hồi bạn cặp."})
    for row in exploratory["findings"]:
        bug_id=row["id"]
        bugs.append({"id":bug_id,"title":row["title"],"category":"Visual" if bug_id=="X03" else "Functional",
                     "severity":"Major" if bug_id=="X02" else "Minor","priority":"High" if bug_id=="X02" else "Medium",
                     "requirement":"Ngoài đặc tả gốc / tiêu chí exploratory UX",
                     "requirement_quote":"Không có yêu cầu tương ứng trong PDF/SRS gốc; expected là tiêu chí usability đề xuất, cần chủ hệ thống xác nhận.",
                     "spec_relation":"Ngoài đặc tả gốc; lỗi khám phá, không cài chủ đích. Không tính vào20lỗi trong Bug Manifest.","seeded":False,
                     "url":"/#requests" if bug_id=="X01" else "Hộp thoại Tạo yêu cầu" if bug_id=="X02" else "/#request/:id (390px)",
                     "preconditions":"TRAINING_BUGS=0; Edge155; database demo tạm. Cùng tab cho X01/X02; viewport390×844 cho X03.",
                     "steps":row["steps"],"expected":row["expected"],"actual":row["actual"],"testcases":[f"TC-EXP-{bug_id[-2:]}"],
                     "evidence":[f"evidence/exploratory/{name}" for name in row["screenshots"]]+["evidence/exploratory/results.json"],
                     "file":"static/app.js" if bug_id!="X03" else "static/style.css",
                     "date":date_of(exploratory["finished_at_utc"]),"environment":"Windows / Edge155.0.4283.62 /390×844; baseline, database tạm.",
                     "status":"Mới","owner_response":"Phát hiện nội bộ có bằng chứng; đề xuất cải thiện UX, chưa có quyết định/phản hồi độc lập của bạn cặp."})
    return bugs


def add_validation(ws,col,start,end,values):
    validation=DataValidation(type="list",formula1='"'+','.join(values)+'"',allow_blank=True)
    validation.errorTitle="Chọn giá trị trong danh sách"
    validation.error="Dùng giá trị theo biểu mẫu."
    validation.showErrorMessage=True
    ws.add_data_validation(validation)
    validation.add(f"{get_column_letter(col)}{start}:{get_column_letter(col)}{end}")


def base_workbook():
    wb=Workbook()
    wb.remove(wb.active)
    wb.properties.creator=IDENTITY
    wb.properties.title="StudentCare – hồ sơ tự kiểm thử – 23110193"
    wb.properties.subject=SCOPE
    wb.properties.description=SOURCE
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    return wb


def titled_sheet(wb,name,title,subtitle,width):
    ws=wb.create_sheet(name)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=width)
    ws.cell(1,1,title).font=Font(name="Calibri",size=16,bold=True,color="FFFFFF")
    ws.cell(1,1).fill=PatternFill("solid",fgColor=NAVY)
    ws.cell(1,1).alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=width)
    ws.cell(2,1,subtitle).font=Font(name="Calibri",size=10,color="53695B",italic=True)
    ws.cell(2,1).alignment=Alignment(wrap_text=True,vertical="top")
    ws.row_dimensions[2].height=32
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape"
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToHeight=0
    ws.page_setup.fitToWidth=1
    ws.print_options.horizontalCentered=True
    return ws


def table(ws,headers,rows,start=4,widths=None):
    for col,header in enumerate(headers,1):
        cell=ws.cell(start,col,header)
        cell.font=Font(name="Calibri",size=11,bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor=GREEN)
        cell.alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[start].height=38
    thin=Side(style="hair",color="DCE5DF")
    for index,row in enumerate(rows,start+1):
        for col,value in enumerate(row,1):
            if isinstance(value,(dict,list)):value=text(value)
            cell=ws.cell(index,col,value)
            cell.font=Font(name="Calibri",size=10,color="23382C")
            cell.fill=PatternFill("solid",fgColor="FFFFFF" if (index-start)%2 else "F3F7F4")
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            cell.border=Border(bottom=thin)
            if value in ("Fail","Blocked"):
                cell.fill=PatternFill("solid",fgColor=RED)
                cell.font=Font(name="Calibri",size=10,bold=True,color="9D3531")
            elif value=="Pass":
                cell.fill=PatternFill("solid",fgColor=LIGHT)
                cell.font=Font(name="Calibri",size=10,bold=True,color=GREEN)
        longest=max((len(str(value or "")) for value in row),default=0)
        ws.row_dimensions[index].height=min(180,max(42,28+longest//70*12))
    for col in range(1,len(headers)+1):
        ws.column_dimensions[get_column_letter(col)].width=(widths[col-1] if widths and col<=len(widths) else 27)
    ws.freeze_panes=f"C{start+1}"
    ws.auto_filter.ref=f"A{start}:{get_column_letter(len(headers))}{start+len(rows)}"
    ws.print_title_rows=f"1:{start}"
    ws.print_area=f"A1:{get_column_letter(len(headers))}{start+len(rows)}"


def link_cell(ws,row,col,paths,self_test=False):
    if not paths:return
    values=[workbook_link(path,self_test) for path in paths]
    cell=ws.cell(row,col,"\n".join(values))
    cell.hyperlink=values[0]
    cell.font=Font(name="Calibri",size=10,color="0563C1",underline="single")
    cell.alignment=Alignment(wrap_text=True,vertical="top")


def manifest(headers,bugs):
    wb=base_workbook()
    ws=titled_sheet(wb,"Bug_Manifest","DANH SÁCH LỖI CÀI SẴN (BUG MANIFEST)",
                   "NỘP RIÊNG GIẢNG VIÊN — không gửi bạn cặp. 20lỗi chủ đích; giữ nguyên sản phẩm đến demo. "+SOURCE,13)
    info=[("Họ và tên",IDENTITY),("MSSV",MSSV),("Tên hệ thống","StudentCare v1.1"),("Link sản phẩm",SITE),
          ("Tester (bạn cặp)",partner_info()),("Ngày lập hồ sơ",datetime.now(LOCAL_TZ).strftime("%d/%m/%Y"))]
    for row,(label,value) in enumerate(info,4):
        ws.cell(row,1,label).font=Font(bold=True,color=NAVY)
        ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=4)
        ws.cell(row,2,value).alignment=Alignment(wrap_text=True)
        ws.cell(row,2).fill=PatternFill("solid",fgColor=YELLOW)
        ws.row_dimensions[row].height=29
    ws.cell(4,6,"Tự kiểm tra số lượng").font=Font(bold=True,color=NAVY)
    for col,value in [(7,"Số lượng"),(8,"Tối thiểu"),(9,"Trạng thái")]:ws.cell(4,col,value).font=Font(bold=True,color=NAVY)
    rows=[]
    for bug in bugs:
        if not bug["seeded"]:continue
        rows.append([bug["id"],bug["category"],bug["severity"],bug["requirement"],bug["url"],bug["requirement_quote"],
                     bug["title"],text(bug["steps"]),bug["expected"],text(bug["actual"]),"\n".join(bug["evidence"]),bug["file"],
                     bug["spec_relation"]+"\n"+SCOPE])
    start=16;end=start+len(rows)
    for row,category,minimum in [(5,"Functional",5),(6,"Visual",3),(7,"Content",3),(8,"Performance",2),(9,"Crash",2)]:
        ws.cell(row,6,category);ws.cell(row,7,f'=COUNTIF(B{start+1}:B{end},F{row})');ws.cell(row,8,minimum)
        ws.cell(row,9,f'=IF(G{row}>=H{row},"Đạt","Chưa đạt")')
    ws.cell(10,6,"Tổng số lỗi");ws.cell(10,7,"=SUM(G5:G9)");ws.cell(10,8,15);ws.cell(10,9,'=IF(G10>=H10,"Đạt","Chưa đạt")')
    ws.cell(11,6,"Mức Critical + Major");ws.cell(11,7,f'=COUNTIF(C{start+1}:C{end},"Critical")+COUNTIF(C{start+1}:C{end},"Major")')
    ws.cell(11,8,3);ws.cell(11,9,'=IF(G11>=H11,"Đạt","Chưa đạt")')
    ws.merge_cells("A13:M13");ws.cell(13,1,"Nguồn yêu cầu: SRS v1.1 của chủ hệ thống; 10B truy vết nghiệp vụ, 10V/C/P/R bám NFR bổ sung. Không tính3bug exploratory vào manifest.").alignment=Alignment(wrap_text=True)
    ws.row_dimensions[13].height=30
    table(ws,headers["Bug_Manifest"],rows,start,[12,17,18,23,43,63,48,65,57,65,55,38,65])
    for i,bug in enumerate((b for b in bugs if b["seeded"]),start+1):link_cell(ws,i,11,bug["evidence"])
    add_validation(ws,2,start+1,end,CATEGORIES);add_validation(ws,3,start+1,end,SEVERITIES)
    path=PRIVATE/f"Mau_02_Bug_Manifest_{MSSV}.xlsx"
    wb.save(path)
    return path


def review_rows():
    return [
        [1,"PDF gốc §9 / SRS v1.1 §6 NFR-01..03","Các chỉ tiêu hiệu năng cụ thể cần đo/duyệt sau.","Draft gốc chưa có ngưỡng, tải, dữ liệu hay percentile; không thể kết luận đạt hiệu năng.","Không kiểm thử được","Cao","Bổ sung ≥3NFR hiệu năng: p95, lỗi%, tải1–2/1→2→4, thời lượng300s/120s, dữ liệu5hồ sơ.","Chấp nhận nội bộ","SRS v1.1 §6 đã ghi ngưỡng/môi trường. Không coi là số đo đã đạt; không có xác nhận bạn cặp."],
        [2,"PDF gốc TBD-07 / SRS FR-13,NFR-09","Loại, dung lượng, số lượng tệp: TBD.","Chưa xác định biên file rỗng/kích thước/định dạng nên thiếu oracle cho BVA.","Thiếu yêu cầu","Cao","Chốt1..5242880byte,≤3tệp,request≤16MiB; whitelist; file rỗng400; rollback nguyên tử.","Chấp nhận nội bộ","Đã cụ thể hóa FR-13/NFR-09; đây là lựa chọn bản thực hành, chưa là chính sách trường."],
        [3,"PDF gốc TBD-02/04/09 / SRS FR-15","Mở lại trực tiếp hay duyệt; cán bộ khi mở lại; REJECTED: TBD.","Cơ chế mở lại và phân công chưa đủ để viết testcase trạng thái nhất quán.","Mơ hồ","Cao","Chỉ COMPLETED mở lại; giữ cán bộ hợp lệ, nếu không về NEW; REJECTED409; ghi REOPEN.","Chấp nhận nội bộ","SRS FR-15/FLOW-03 chọn mở lại trực tiếp, không thêm REOPEN_REQUESTED; cần bạn cặp review."],
        [4,"PDF gốc F05 / SRS §8","Lỗi ghi thông báo không được làm mất kết quả nghiệp vụ.","Cài đặt dùng cùng transaction, lỗi notification sẽ rollback nghiệp vụ; khác đề xuất PDF.","Không nhất quán","Cao","Ghi rõ giới hạn; muốn đáp ứng đầy đủ cần outbox/retry riêng và testcase fault injection.","Cần xác nhận thêm","SRS §8 nêu khác biệt thực tế; chưa triển khai outbox. Không đánh dấu đã sửa mã hoặc đã được bạn cặp chấp thuận."],
        [5,"PDF gốc §9 Retry / SRS §7,NFR-07,API v1.1","Chống tạo trùng yêu cầu và thay đổi trạng thái khi gửi lại.","Draft chưa định nghĩa phạm vi key, cùng key nhưng body khác và retry multipart khác boundary.","Thiếu chi tiết","Trung bình","Key theo user+method+path; cùng body trả kết quả cũ, khác body409; hash nội dung tệp không phụ thuộc boundary.","Chấp nhận nội bộ","Bổ sung API và model Idempotency; đã có testcase JSON/multipart. Chưa có review độc lập."],
    ]


def information(wb,functional,nfr):
    ws=titled_sheet(wb,"Thong_tin","BỘ BIỂU MẪU KIỂM THỬ – HỒ SƠ TỰ KIỂM THỬ",SCOPE,6)
    entries=[("Tester – Họ và tên",IDENTITY),("Tester – MSSV",MSSV),("Hệ thống được kiểm thử","StudentCare v1.1 — hệ thống của chính sinh viên"),
             ("Chủ hệ thống – Họ tên, MSSV",f"{IDENTITY} – {MSSV}"),("Link sản phẩm",SITE),("Link GitHub",REPO),
             ("Phiên bản SRS dùng làm căn cứ","SRS v1.1 ngày24/10/2026; requirements bổ sung được ghi rõ"),
             ("Thời gian kiểm thử (từ – đến)","Theo timestamp trong JSON; hiển thị Test_Log theo Asia/Saigon"),
             ("Bạn cặp / kiểm thử chéo",partner_info()),("Trạng thái phần tester","CHƯA HOÀN THÀNH ĐIỀU KIỆN KIỂM THỬ CHÉO. File này chỉ chứa self-test thực chạy."),
             ("Nguồn mẫu",DRIVE),("Cách sử dụng mẫu",SOURCE)]
    for row,(label,value) in enumerate(entries,4):
        ws.cell(row,1,label).font=Font(bold=True,color=NAVY)
        ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=6)
        cell=ws.cell(row,2,value);cell.alignment=Alignment(wrap_text=True,vertical="top");cell.fill=PatternFill("solid",fgColor=YELLOW)
        ws.row_dimensions[row].height=35
    ws.cell(18,1,"CÁC SHEET TRONG FILE").font=Font(bold=True,color=NAVY)
    for row,(name,description) in enumerate([("Review_SRS","5 vấn đề review nội bộ; không giả là review của bạn cặp."),("TC_Chuc_nang",f"{len(functional)} testcase, ≥3/FR,9 Use case / E2E; actual và evidence thật."),
                    ("TC_Phi_chuc_nang",f"{len(nfr)} testcase phủ NFR-01..10; chưa chạy giữ Chưa chạy."),("Test_Log","Log thực thi có nguồn timestamp và môi trường cục bộ."),("Bug_Report","23bug:20cài sẵn+3ngoài đặc tả. Không chia sẻ cho bạn cặp khi chưa được phép."),("Tong_hop","Công thức tự đếm, cần Excel tính lại khi mở; phạm vi chỉ self-test.")],19):
        ws.cell(row,1,name).font=Font(bold=True,color=GREEN);ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=6);ws.cell(row,2,description).alignment=Alignment(wrap_text=True)
        ws.row_dimensions[row].height=30
    ws.column_dimensions["A"].width=40
    for col in "BCDEF":ws.column_dimensions[col].width=20
    ws.freeze_panes="B4"
    ws.print_area="A1:F25"


def log_rows(functional_report,nfr_report,extra,exploratory,tooling):
    entries=[]
    def append(start,finish,version,environment,scope,total,passed,failed,blocked,bugs,note):
        if not start:return
        began=datetime.fromisoformat(start.replace("Z","+00:00")).astimezone(LOCAL_TZ)
        ended=datetime.fromisoformat((finish or start).replace("Z","+00:00")).astimezone(LOCAL_TZ)
        entries.append([len(entries)+1,began.strftime("%d/%m/%Y"),began.strftime("%H:%M:%S")+" – "+ended.strftime("%H:%M:%S"),
                        version,environment,scope,total,passed,failed,blocked,", ".join(bugs),note])
    cases=functional_report["cases"]
    beginning=min(r["executed_at"] for r in cases)
    append(beginning,functional_report["executed_at"],"StudentCare local/training; source hashes trong results.json","Flask test_client / SQLite tạm, mỗi case DB độc lập",
           f"TC-F-001..088; TC-FLOW01-01..TC-FLOW03-03",len(cases),functional_report["counts"]["PASS"],functional_report["counts"]["FAIL"],functional_report["counts"]["ERROR"],
           sorted({r["bug_id"] for r in cases if r.get("bug_id") and r["status"]=="FAIL"}),"Self-test, không phải bằng chứng bạn cặp; evidence/functional/training/results.json.")
    extra_cases=[r for r in extra["results"] if r["mode"]=="training"]
    append(extra["started_at_utc"],extra["finished_at_utc"],"StudentCare local/training / hash nguồn ghi trong results.json","Edge155 / HTTP tuần tự1người dùng",
           "NFR-01/02/04/05/08: V01..V03,C01..C03,P01/P02,R01/R02",len(extra_cases),sum(r["status"]=="PASS" for r in extra_cases),sum(r["status"]=="FAIL" for r in extra_cases),0,
           [r["id"] for r in extra_cases if r["status"]=="FAIL"],"Không gọi5mẫu tuần tự là kiểm thử tải. evidence/additional/results.json.")
    append(exploratory["started_at_utc"],exploratory["finished_at_utc"],"StudentCare local/baseline; TRAINING_BUGS=0","Edge155 /390×844 / SQLite tạm",
           "TC-EXP-01..03: X01/X02/X03",3,0,3,0,["X01","X02","X03"],"3phát hiện ngoài đặc tả; chưa có phản hồi bạn cặp. evidence/exploratory/results.json.")
    for name,run in tooling.get("runs",{}).items():
        if name not in {"newman","puppeteer","jmeter-load","jmeter-step"}:continue
        append(run.get("started"),run.get("finished"),"StudentCare local/training",tooling.get("environment",{}).get("server","Cục bộ"),name,
               None,None,None,None,[],f"Tool exit={run.get('exit_code')}; dùng báo cáo gốc để đọc assertion/metrics, không coi exitcode0 là mọi TC đạt. {run.get('log','')}")
    return entries


def summary_sheet(wb,fc_count,nfr_count,bug_count,review_count):
    ws=titled_sheet(wb,"Tong_hop","TỔNG HỢP SỐ LIỆU (TỰ ĐỘNG)","Chỉ tổng hợp hồ sơ tự kiểm thử. Excel sẽ tính lại công thức khi mở; chưa thể thay cho kiểm thử chéo.",6)
    f_end=fc_count+4;n_end=nfr_count+4;b_end=bug_count+4;r_end=review_count+4
    ws.append([])
    for cell,value in {"A4":"1. Kiểm tra số lượng tối thiểu","B4":"Hiện có","C4":"Tối thiểu","D4":"Trạng thái"}.items():ws[cell]=value
    for row,label,formula,minimum in [
        (5,"Test case chức năng",f"=COUNTA(TC_Chuc_nang!A5:A{f_end})",60),
        (6,"Test case phi chức năng",f"=COUNTA(TC_Phi_chuc_nang!A5:A{n_end})",10),
        (7,"Test case luồng E2E (Use case / E2E)",f'=COUNTIF(TC_Chuc_nang!C5:C{f_end},"Use case / E2E")',9),
        (8,"Số kỹ thuật thiết kế đã dùng",'+'.join(f'IF(COUNTIF(TC_Chuc_nang!C5:C{f_end},"{name}")>0,1,0)' for name in TECHNIQUES),4)]:
        ws.cell(row,1,label);ws.cell(row,2,formula if formula.startswith("=") else "="+formula);ws.cell(row,3,minimum);ws.cell(row,4,f'=IF(B{row}>=C{row},"Đạt","Chưa đạt")')
    for col,val in enumerate(["2. Kết quả thực thi","Chức năng","Phi chức năng","Tổng"],1):ws.cell(11,col,val)
    labels=["Tổng số test case","Pass","Fail","Blocked","Chưa chạy","Đã thực thi","Tỉ lệ Pass / đã thực thi"]
    for row,label in enumerate(labels,12):ws.cell(row,1,label)
    ws["B12"]=f"=COUNTA(TC_Chuc_nang!A5:A{f_end})";ws["C12"]=f"=COUNTA(TC_Phi_chuc_nang!A5:A{n_end})";ws["D12"]="=SUM(B12:C12)"
    for row,status in enumerate(STATUSES,13):
        ws.cell(row,2,f'=COUNTIF(TC_Chuc_nang!J5:J{f_end},"{status}")');ws.cell(row,3,f'=COUNTIF(TC_Phi_chuc_nang!J5:J{n_end},"{status}")');ws.cell(row,4,f"=SUM(B{row}:C{row})")
    for col in "BCD":
        ws[f"{col}17"]=f"=SUM({col}13:{col}14)";ws[f"{col}18"]=f"=IFERROR({col}13/{col}17,0)";ws[f"{col}18"].number_format="0.0%"
    for col,label in enumerate(["3. Bug theo loại lỗi",*SEVERITIES,"Tổng"],1):ws.cell(21,col,label)
    for row,category in enumerate(CATEGORIES,22):
        ws.cell(row,1,category)
        for col,severity in enumerate(SEVERITIES,2):ws.cell(row,col,f'=COUNTIFS(Bug_Report!C5:C{b_end},$A{row},Bug_Report!D5:D{b_end},{get_column_letter(col)}$21)')
        ws.cell(row,6,f"=SUM(B{row}:E{row})")
    ws["A27"]="Tổng"
    for col in "BCDEF":ws[f"{col}27"]=f"=SUM({col}22:{col}26)"
    ws["A30"]="4. Bug theo trạng thái phản hồi";ws["B30"]="Số lượng"
    for row,status in enumerate(BUG_STATES,31):ws.cell(row,1,status);ws.cell(row,2,f'=COUNTIF(Bug_Report!P5:P{b_end},A{row})')
    ws["A38"]="5. Test case theo kỹ thuật";ws["B38"]="Số lượng"
    for row,name in enumerate(TECHNIQUES,39):ws.cell(row,1,name);ws.cell(row,2,f'=COUNTIF(TC_Chuc_nang!C5:C{f_end},A{row})')
    ws["A49"]="6. Review SRS";ws["B49"]="Số lượng";ws["A50"]="Tổng số vấn đề đã nêu";ws["B50"]=f"=COUNTA(Review_SRS!A5:A{r_end})"
    ws["A51"]="Được chủ hệ thống chấp nhận nội bộ";ws["B51"]=f'=COUNTIF(Review_SRS!H5:H{r_end},"Chấp nhận nội bộ")'
    ws["A53"]="Kiểm thử chéo với bạn cặp";ws["B53"]="CHƯA XÁC NHẬN"
    ws["A54"]="Kết luận sản phẩm";ws["B54"]='=IF(SUM(B14:C14)>0,"CHƯA ĐẠT NGHIỆM THU","CẦN REVIEW")'
    for row in [4,11,21,30,38,49]:
        for cell in ws[row][:6]:cell.fill=PatternFill("solid",fgColor=GREEN);cell.font=Font(bold=True,color="FFFFFF");cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[row].height=30
    ws.column_dimensions["A"].width=55
    for col in "BCDEF":ws.column_dimensions[col].width=23
    for row in ws.iter_rows(min_row=5,max_row=54):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    ws.freeze_panes="B5";ws.print_area="A1:F54"


def test_workbook(headers,functional,nfr,bugs):
    wb=base_workbook()
    information(wb,functional,nfr)
    reviews=review_rows()
    ws=titled_sheet(wb,"Review_SRS","PHIẾU REVIEW SRS NỘI BỘ",SCOPE+" Căn cứ PDF Draft gốc và SRS v1.1; không ghi giả phản hồi bạn cặp.",9)
    table(ws,headers["Review_SRS"],reviews,widths=[8,35,52,58,27,18,65,30,65])
    ws=titled_sheet(wb,"TC_Chuc_nang","TEST CASE CHỨC NĂNG",f"97testcase tự thực thi, gồm9Use case / E2E. "+SCOPE,13)
    functional_rows=[]
    for row in functional:
        note=f"{SCOPE}\nMôi trường: Flask test_client/SQLite tạm; TRAINING_BUGS=1.\nBaseline={row.get('baseline_status','chưa có')}; evidence={relative_evidence(row['evidence'])}"
        designed_steps=row["steps"]
        # Generic static planning instructions are supplemented by exact executed operations.
        if "Gọi API với dữ liệu của testcase." in designed_steps:
            designed_steps=row.get("executed_steps",designed_steps)
        functional_rows.append([row["id"],row["function_id"]+" – "+FUNCTIONS.get(row["function_id"],"")+(" / "+E2E_SRS_FLOWS[row["id"]] if row["id"] in E2E_SRS_FLOWS else ""),
                                technique(row),row["title"],row["preconditions"],text(row["data"]),text(designed_steps),row["expected"],"High" if row.get("bug_id") or row.get("test_kind")=="Permission" else "Medium",
                                result_status(row["status"]),row["actual"] if row["status"]!="PASS" else "",row.get("bug_id","") if row["status"]!="PASS" else "",note])
    table(ws,headers["TC_Chuc_nang"],functional_rows,widths=[20,33,28,47,48,66,75,60,15,15,65,15,73])
    add_validation(ws,3,5,4+len(functional),TECHNIQUES);add_validation(ws,9,5,4+len(functional),["High","Medium","Low"]);add_validation(ws,10,5,4+len(functional),STATUSES)
    for row_index,row in enumerate(functional,5):
        ws.cell(row_index,13).hyperlink=workbook_link(row["evidence"],True)
    ws=titled_sheet(wb,"TC_Phi_chuc_nang","TEST CASE PHI CHỨC NĂNG",SCOPE+" Đo trên bản sao cục bộ; ngưỡng SRS là mục tiêu, không mặc nhiên đạt.",12)
    nfr_rows=[]
    for row in nfr:
        evidence=relative_evidence(row.get("evidence",""))
        nfr_rows.append([row["id"],row["requirement"],NFR_GROUPS.get(row["requirement"],"Phi chức năng"),row["title"],row["method"],
                         "Windows / local isolated SQLite / training=1; điều kiện chi tiết trong evidence, không phải VPS.",text(row.get("steps",[])),row["expected"],text(row["actual"]),
                         result_status(row["status"]),row.get("bug_id",""),SCOPE+"\n"+evidence])
    table(ws,headers["TC_Phi_chuc_nang"],nfr_rows,widths=[22,17,25,50,35,60,68,62,70,16,15,60])
    add_validation(ws,10,5,4+len(nfr),STATUSES)
    for index,row in enumerate(nfr,5):ws.cell(index,12).hyperlink=workbook_link(row["evidence"],True)
    functional_report=read(EVIDENCE/"functional"/"training"/"results.json")
    nfr_report=read(EVIDENCE/"nonfunctional"/"results.json")
    extra=read(EVIDENCE/"additional"/"results.json")
    exploratory=read(EVIDENCE/"exploratory"/"results.json")
    tooling=read(EVIDENCE/"tooling"/"summary.json")
    ws=titled_sheet(wb,"Test_Log","NHẬT KÝ THỰC THI",SCOPE+" Thời gian chuyển từ timestamp UTC sang Asia/Saigon; không tạo phiên test giả.",12)
    logs=log_rows(functional_report,nfr_report,extra,exploratory,tooling)
    table(ws,headers["Test_Log"],logs,widths=[8,16,26,57,55,65,16,14,14,14,47,80])
    ws=titled_sheet(wb,"Bug_Report","DANH SÁCH BUG REPORT – NỘI BỘ",SCOPE+" Có23lỗi thực tái hiện,20cài sẵn và3khám phá; không gửi bạn cặp như manifest.",17)
    bug_rows=[]
    for bug in bugs:
        bug_rows.append([bug["id"],bug["title"],bug["category"],bug["severity"],bug["priority"],bug["requirement"],", ".join(bug["testcases"]),
                         SITE+" · "+bug["url"]+" (evidence đo cục bộ)",bug["environment"],bug["preconditions"],text(bug["steps"]),text(bug["actual"]),
                         bug["requirement_quote"]+"\nExpected: "+bug["expected"],"\n".join(bug["evidence"]),bug["date"],bug["status"],bug["owner_response"]])
    table(ws,headers["Bug_Report"],bug_rows,widths=[12,58,19,20,17,36,48,60,60,65,75,75,75,60,18,25,75])
    for i,bug in enumerate(bugs,5):link_cell(ws,i,14,bug["evidence"],True)
    for col,values in [(3,CATEGORIES),(4,SEVERITIES),(5,["High","Medium","Low"]),(16,BUG_STATES)]:add_validation(ws,col,5,4+len(bugs),values)
    summary_sheet(wb,len(functional),len(nfr),len(bugs),len(reviews))
    path=SELF/f"Mau_03_Bo_bieu_mau_TU_KIEM_THU_{MSSV}.xlsx"
    wb.save(path)
    return path


def validate(manifest_path,test_path,functional,nfr,bugs,headers):
    manifest_wb=load_workbook(manifest_path,data_only=False)
    test_wb=load_workbook(test_path,data_only=False)
    expected_sheets=["Thong_tin","Review_SRS","TC_Chuc_nang","TC_Phi_chuc_nang","Test_Log","Bug_Report","Tong_hop"]
    assert test_wb.sheetnames==expected_sheets,test_wb.sheetnames
    assert [c.value for c in manifest_wb["Bug_Manifest"][16]]==headers["Bug_Manifest"]
    for name in expected_sheets[1:-1]:assert [c.value for c in test_wb[name][4]]==headers[name],name
    assert test_wb["TC_Chuc_nang"].max_row==len(functional)+4
    assert test_wb["TC_Phi_chuc_nang"].max_row==len(nfr)+4
    assert test_wb["Bug_Report"].max_row==len(bugs)+4
    assert manifest_wb["Bug_Manifest"].max_row==20+16
    assert len(functional)>=60 and len(nfr)>=10
    assert {row["function_id"] for row in functional}==set(FUNCTIONS)
    assert {row["requirement"] for row in nfr}==set(NFR_GROUPS)
    assert sum(row.get("technique")=="Use case / E2E" for row in functional)>=9
    assert len({technique(row) for row in functional})>=4
    assert all(value>=3 for value in Counter(row["function_id"] for row in functional).values())
    assert len([b for b in bugs if not b["seeded"]])==3
    seeded_counts=Counter(b["category"] for b in bugs if b["seeded"])
    assert seeded_counts=={"Functional":10,"Visual":3,"Content":3,"Performance":2,"Crash":2}
    assert sum(b["severity"] in {"Critical","Major"} for b in bugs if b["seeded"])>=3
    missing=[]
    for bug in bugs:
        for item in bug["evidence"]:
            if not (SUB/relative_evidence(item)).exists():missing.append(item)
    assert not missing,missing
    summary={"generated_at":datetime.now(LOCAL_TZ).isoformat(),"scope":SCOPE,"template":SOURCE,
             "functional":len(functional),"functional_counts":dict(Counter(r["status"] for r in functional)),"e2e":9,
             "techniques":dict(Counter(technique(r) for r in functional)),"nfr":len(nfr),"nfr_counts":dict(Counter(r["status"] for r in nfr)),
             "seeded_bugs":20,"seeded_categories":dict(seeded_counts),"exploratory_bugs":3,"bug_report_total":len(bugs),
             "formula_validation":"Formula strings and referenced sheet/range structure checked; Excel recalculates on opening. No cached values fabricated.",
             "paths":[str(manifest_path.relative_to(ROOT)),str(test_path.relative_to(ROOT))]}
    (SUB/"data"/"workbook_validation.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    return summary


def main():
    PRIVATE.mkdir(parents=True,exist_ok=True);SELF.mkdir(parents=True,exist_ok=True)
    template_file=ROOT/"test-results"/"template-reading"/"sheet-headers.json"
    headers=read(template_file) if template_file.exists() else TEMPLATE_HEADERS
    assert headers==TEMPLATE_HEADERS,"Template headers changed; review before rebuilding."
    functional=read(SUB/"data"/"functional_cases.json")
    nfr=read(EVIDENCE/"nonfunctional"/"results.json")["cases"]
    bugs=build_catalog(functional,nfr)
    (SUB/"data"/"bug_catalog.json").write_text(json.dumps(bugs,ensure_ascii=False,indent=2),encoding="utf-8")
    manifest_path=manifest(headers,bugs)
    test_path=test_workbook(headers,functional,nfr,bugs)
    summary=validate(manifest_path,test_path,functional,nfr,bugs,headers)
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
