"""Package two disclosed, AI-assisted coursework dossiers using captured evidence."""
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path

from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[1]
PEER=ROOT.parent/'partner-shiloskin'
SUB=ROOT/'submission';PAIR=SUB/'pair'
sys.path.insert(0,str(ROOT/'tools'))
import build_submission as writer
import build_workbooks as template
HEADERS=template.TEMPLATE_HEADERS
DISCLOSURE='Codex tự động thực thi theo yêu cầu của người dùng; cả hai ứng dụng được AI hỗ trợ xây dựng. Đây không phải kiểm thử mù độc lập; không giả thao tác tay, chữ ký hay điểm chấm.'

def read(path,default=None):
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else default
def dump(path,value):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
def text(value):
    return value if isinstance(value,str) else json.dumps(value,ensure_ascii=False,indent=2)
def copy(source,target):
    if source.is_dir():shutil.copytree(source,target,dirs_exist_ok=True,ignore=shutil.ignore_patterns('node_modules','__pycache__','.git','aborted-*'))
    elif source.exists():target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(source,target)
def document(folder,stem,title,body,name,msv,project,base):
    folder.mkdir(parents=True,exist_ok=True)
    writer.IDENTITY=f'{name} – {msv}';writer.PROJECT=project;writer.BASE=base
    writer.write_document(stem,title,body,folder)

def normalize_tooling():
    report=read(PEER/'evidence/tooling/summary.json',{})
    for kind in ('load','step'):
        run=report.get('runs',{}).get('jmeter-'+kind)
        if not run or not run.get('duration_verified'):continue
        source=PEER/Path(run['raw_results']).parent
        canonical=PEER/'evidence/tooling'/('jmeter-'+kind)
        if source.resolve()!=canonical.resolve():copy(source,canonical)
    if all('jmeter-'+key in report.get('runs',{}) for key in ['load','step']):
        subprocess.run([sys.executable,'-X','utf8',str(PEER/'tooling/summarize_jmeter.py')],check=True,stdout=subprocess.DEVNULL)
    copy(PEER/'evidence/tooling',SUB/'partner/evidence/tooling-final')

def shilo_catalog():
    base=SUB/'partner';functional=read(base/'data/functional_cases.json',[])
    proof=read(base/'evidence/seeded_functional/results.json',{}).get('results',[])
    extra=read(base/'evidence/additional/results.json',{}).get('results',[])
    exploratory=read(base/'evidence/exploratory/results.json',{}).get('findings',[])
    mapping={'S01':'FR-05','S02':'FR-05','S03':'FR-07','S04':'FR-07','S05':'FR-06','S06':'FR-09','S07':'FR-15','S08':'FR-13','S09':'FR-16','S10':'FR-11'}
    catalog=[]
    for row in proof+[r for r in extra if r.get('mode')=='training']+exploratory:
        bug=row['id'];seed=not bug.startswith('X');folder='seeded_functional' if bug.startswith('S') else 'exploratory' if not seed else 'additional'
        files=row.get('evidence_files',row.get('screenshots',[]))
        if not files:files=[p.name for p in (base/'evidence'/folder).glob(bug+'*.png')]
        files=['evidence/'+folder+'/'+f for f in files]
        files+=['evidence/'+folder+'/results.json']
        catalog.append({'id':bug,'title':row['title'],'category':row.get('category','Visual' if bug=='X03' else 'Functional'),
            'severity':'Major' if bug.startswith(('S','P','R')) else 'Minor','priority':'High' if bug.startswith(('S','P','R')) else 'Medium',
            'requirement':mapping.get(bug,row.get('requirement','Exploratory ngoài đặc tả')),'requirement_quote':row['expected'],
            'seeded':seed,'spec_relation':'SRS ShiloSkin v1.0' if seed else 'Khám phá ngoài đặc tả gốc, heuristic UX',
            'url':row.get('url','https://sv21.bcse-vju.com'),'preconditions':'Demo database riêng; tài khoản mẫu đúng vai trò; chế độ nêu trong raw evidence.',
            'steps':row.get('steps',[]),'expected':row['expected'],'actual':row['actual'],'evidence':files,
            'testcases':[r['id'] for r in functional if r.get('bug_id')==bug],
            'file':'app.py / training_extras.py / static/training.*','date':'25/10/2026','environment':'Windows, Edge155, SQLite tạm, không phải số đo VPS',
            'status':'Đã xác nhận','owner_response':'Giữ nguyên lỗi cho demo; xác nhận bằng automation, sinh viên cần rà soát.'})
    assert len(catalog)==23,(len(catalog),[r['id'] for r in catalog])
    dump(base/'data/bug_catalog.json',catalog);return catalog

def shilo_nfr(functional):
    base=SUB/'partner';rows=[]
    for r in read(base/'evidence/additional/results.json')['results']:
        if r.get('mode')=='training':rows.append({'id':'TC-'+r['id'],'requirement':r['requirement'],'title':r['title'],'method':'HTTP/Edge DOM và ảnh thực tế','steps':r.get('steps',[]),'expected':r['expected'],'actual':r['actual'],'status':r['status'],'bug_id':r['id'],'evidence':'evidence/additional/results.json'})
    guardlog=base/'evidence/nonfunctional/backend-guards.txt';guardlog.parent.mkdir(parents=True,exist_ok=True)
    result=subprocess.run([sys.executable,'-X','utf8','-m','unittest','discover','-s','tests','-v'],cwd=PEER,capture_output=True,text=True,encoding='utf-8',env={**os.environ,'TRAINING_BUGS':'0'})
    guardlog.write_text(result.stdout+result.stderr,encoding='utf-8')
    assert result.returncode==0,'Backend guards failed; inspect raw log'
    for nfr,title in [('NFR-06','RBAC, CSRF, thu hồi phiên và riêng tư'),('NFR-07','Giao dịch nguyên tử, hai người mua và retry')]:
        rows.append({'id':'TC-'+nfr,'requirement':nfr,'title':title,'method':'unittest guard thực chạy','steps':['Chạy python -m unittest discover -s tests -v.','Đối chiếu log của các guard tương ứng.'],'expected':'Không trường hợp trái quyền được xử lý; đúng một lần chuyển quyền/số dư; guard không lỗi.','actual':result.stderr,'status':'PASS','bug_id':'','evidence':'evidence/nonfunctional/backend-guards.txt'})
    for nfr,bug,title in [('NFR-09','S07','Float phải nằm trong [0,1]'),('NFR-10','S10','Mua thành công phải ghi BUY audit')]:
        matches=[r for r in functional if r.get('bug_id')==bug];assert matches
        row=matches[0];rows.append({'id':'TC-'+nfr,'requirement':nfr,'title':title,'method':'Assertion trong functional trace','steps':row['steps'],'expected':row['expected'],'actual':row['actual'],'status':row['status'],'bug_id':bug,'evidence':row['evidence']})
    comparisons=read(PEER/'evidence/tooling/nfr-comparison.json',{}).get('checks',[])
    for requirement in ('NFR-01','NFR-02'):
        measured=[r for r in comparisons if r['requirement']==requirement]
        if measured:
            row=next(r for r in rows if r['requirement']==requirement);row['actual']={'JMeter':measured,'pilot':row['actual']};row['method']='JMeter load2users300s + pilot HTTP';row['evidence']='evidence/tooling-final/nfr-comparison.json'
    measured=[r for r in comparisons if r['requirement']=='NFR-03']
    rows.append({'id':'TC-NFR-03','requirement':'NFR-03','title':'Danh sách theo tải bậc1/2/4users','method':'Apache JMeter5.6.3','steps':['Chạy step.jmx đủ120giây/bậc.','Tính P95 từ JTL, đối chiếu SRS.'],'expected':'P95 List/Types≤1000ms,error<1% cho từng bậc.','actual':measured or 'Đang chờ kết quả cuối','status':('PASS' if all(r['result']=='Pass' for r in measured) else 'FAIL') if measured else 'NOT_RUN','bug_id':'','evidence':'evidence/tooling-final/nfr-comparison.json'})
    report={'scope':DISCLOSURE,'executed_at':datetime.now(timezone.utc).isoformat(),'total':len(rows),'counts':dict(Counter(r['status'] for r in rows)),'cases':rows}
    dump(base/'evidence/nonfunctional/results.json',report);return report

def workbook(path,owner,tester,app,functional,nfr,bugs,manifest=False):
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(fullCalcOnLoad=True)
    def sheet(name,headers,rows):
        ws=wb.create_sheet(name);ws.append([app+' · '+owner+' · '+tester]);ws.append([DISCLOSURE]);ws.append([]);ws.append(headers)
        for row in rows:ws.append([text(v) for v in row])
        ws.freeze_panes='B5';ws.auto_filter.ref=f'A4:{get_column_letter(len(headers))}{ws.max_row}'
        for cell in ws[4]:cell.fill=PatternFill('solid',fgColor='176D53');cell.font=Font(color='FFFFFF',bold=True)
        for column in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(column)].width=22 if column<4 else 45
        for row in ws.iter_rows():
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;return ws
    if manifest:
        selected=[b for b in bugs if b['seeded']]
        sheet('Bug_Manifest',HEADERS['Bug_Manifest'],[[b['id'],b['category'],b['severity'],b['requirement'],b['url'],b['requirement_quote'],b['title'],b['steps'],b['expected'],b['actual'],'\n'.join(b['evidence']),b['file'],DISCLOSURE] for b in selected])
        ws=wb.create_sheet('Kiem_tra_so_luong');ws.append(['Loại','Hiện có','Tối thiểu','Kết quả'])
        for i,(kind,minimum) in enumerate([('Functional',5),('Visual',3),('Content',3),('Performance',2),('Crash',2)],2):ws.append([kind,f'=COUNTIF(Bug_Manifest!B5:B{len(selected)+4},A{i})',minimum,f'=IF(B{i}>=C{i},"Đạt","Chưa đạt")'])
    else:
        sheet('Thong_tin',['Trường','Giá trị'],[['Chủ hệ thống',owner],['Tester phụ trách',tester],['Thực thi',DISCLOSURE],['Phiên bản','SRS1.1 StudentCare / SRS1.0 ShiloSkin'],['Kết luận','NO-GO nghiệp vụ: giữ lỗi để demo; chữ ký chưa xác nhận']])
        sheet('Review_SRS',HEADERS['Review_SRS'],[[i,loc,quote,issue,'Ràng buộc/khả kiểm thử','Cao',proposal,'Chờ sinh viên xác nhận','Bản yêu cầu hiện đã cụ thể hóa; không giả review mù'] for i,(loc,quote,issue,proposal) in enumerate([('Dữ liệu','Các trường hợp lệ','Yêu cầu ban đầu chưa có các biên số đo.','Ghi rõ độ dài, float/giá/file tùy hệ thống.'),('Quyền','Ba vai trò','Cần kiểm tra server chứ không chỉ nút UI.','Ma trận quyền×hành động×trạng thái.'),('Trạng thái','Luồng chính/thay thế','Cần chặn thao tác sau kết thúc và cạnh tranh.','State table, retry, hai actor đồng thời.'),('NFR01–03','Phản hồi nhanh','Không thể đo nếu không có tải/dataset/ngưỡng.','Đặt P95, error%, users, ramp và duration.'),('Đo hiệu năng','Số đo môi trường','Số localhost không đại diện VPS.','Ghi môi trường, giữ JTL và giới hạn suy rộng.')],1)])
        sheet('TC_Chuc_nang',HEADERS['TC_Chuc_nang'],[[r['id'],r['function_id'],r.get('technique',template.technique(r)),r['title'],r['preconditions'],r['data'],r.get('executed_steps',r['steps']),r['expected'],'High',r['status'],r['actual'] if r['status']!='PASS' else '',r.get('bug_id',''),r['evidence']] for r in functional])
        sheet('TC_Phi_chuc_nang',HEADERS['TC_Phi_chuc_nang'],[[r['id'],r['requirement'],'NFR',r['title'],r['method'],'SQLite tạm/Windows; raw evidence ghi điều kiện',r.get('steps',[]),r['expected'],r['actual'],r['status'],r.get('bug_id',''),r['evidence']] for r in nfr['cases']])
        counts=Counter(r['status'] for r in functional);nc=nfr['counts']
        sheet('Test_Log',HEADERS['Test_Log'],[[1,'25/10/2026','Xem timestamp mỗi JSON','Bản sao ứng dụng / hash tại evidence','Flask test_client','Functional + 9 E2E',len(functional),counts['PASS'],counts['FAIL'],counts['ERROR'],'Xem Bug_Report','Thực thi bằng Codex, không giả chạy tay'],[2,'25/10/2026','Xem timestamp raw tool','Local HTTP / Edge / JMeter','Windows','NFR + API/UI/load',nfr['total'],nc.get('PASS',0),nc.get('FAIL',0),nc.get('NOT_RUN',0),'Xem Bug_Report','Các pilot lỗi môi trường được giữ riêng, không dùng làm số đo cuối']])
        sheet('Bug_Report',HEADERS['Bug_Report'],[[b['id'],b['title'],b['category'],b['severity'],b['priority'],b['requirement'],', '.join(b['testcases']),b['url'],b['environment'],b['preconditions'],b['steps'],b['actual'],b['expected'],'\n'.join(b['evidence']),b['date'],b['status'],b['owner_response']] for b in bugs])
        ws=wb.create_sheet('Tong_hop');ws.append(['Chỉ số','Chức năng','Phi chức năng','Tổng']);ws.append(['Tổng',len(functional),nfr['total'],'=SUM(B2:C2)'])
        for i,status in enumerate(['PASS','FAIL','ERROR','NOT_RUN'],3):ws.append([status,f'=COUNTIF(TC_Chuc_nang!J5:J{len(functional)+4},A{i})',f'=COUNTIF(TC_Phi_chuc_nang!J5:J{nfr["total"]+4},A{i})',f'=SUM(B{i}:C{i})'])
        ws.append(['E2E',sum(r.get('technique')=='Use case / E2E' for r in functional)]);ws.append(['Kỹ thuật',len({r.get('technique',template.technique(r)) for r in functional})]);ws.append(['Bug',len(bugs)]);ws.append(['Kết luận','NO-GO nghiệm thu; bản thực hành giữ lỗi'])
    for ws in wb:
        for row in ws.iter_rows(min_row=5):
            for cell in row:
                if isinstance(cell.value,str) and cell.value.startswith('evidence/') and '\n' not in cell.value and ';' not in cell.value:cell.hyperlink=cell.value
    path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)

def tester_docs(folder,person,msv,owner,project,functional,nfr,bugs,toolroot):
    fc=Counter(r['status'] for r in functional);nc=nfr['counts']
    newman=read(toolroot/'newman/summary.json',{});pupp=read(toolroot/'puppeteer/summary.json',{});perf=read(toolroot/'nfr-comparison.json',{}).get('checks',[])
    intro=f'Tester phụ trách: {person} – {msv}. Chủ hệ thống: {owner}. Hệ thống được kiểm thử: {project}. Ngày lập25/10/2026.\n\n{DISCLOSURE}\n'
    plan=intro+'''## 1. Giới thiệu
Kiểm chứng SRS, xác định sai lệch chức năng/quyền/trạng thái và NFR. Tham chiếu URD, SRS, API và workbook đi kèm. Bản thực hành giữ lỗi để demo; không kết luận Go khi còn lỗi nghiêm trọng.
## 2. Phạm vi
Trong phạm vi:16FR,3luồng,10NFR; dữ liệu, quyền, phiên, retry, cạnh tranh, giao diện/nội dung. Ngoài phạm vi: dịch vụ thật ngoài ứng dụng; thanh toán/Steam/SSO; tải hạ tầng trường khi chưa đăng ký; năng lực vượt4người dùng.
## 3. Chiến lược
EP, BVA, bảng quyết định, chuyển trạng thái,9Usecase/E2E. API: Postman/Newman,3iterationCSV, request âm và assertion. Puppeteer≥2luồng có ảnh. JMeter: load2users300s/ramp10s; step1→2→4,120s/bậc/ramp5s; cap180sample/phút nhóm. Phân loại Functional/Visual/Content/Performance/Crash; severityCritical/Major/Minor/Trivial theo ảnh hưởng, không đánh đồng với priority.
## 4. Môi trường
Chạy tự động trên bản sao SQLite tạm, Windows/Python3.11,Edge155,Node24,Newman6.2.1,Puppeteer24.38,JMeter5.6.3/JDK21. Số đo localhost không đại diện VPS. Tài khoản và fixture theo biên bản/API; không dùng tiền thật hoặc dữ liệu người dùng thật.
## 5. Tiêu chí
Bắt đầu: SRS có biên kiểm thử, đủ vai trò/dataset, hệ thống truy cập được. Kết thúc: mọi testcase có kết quả/bằng chứng và summary; Go nghiệp vụ cần0Critical/Major và mọiNFRđạt. Tạm dừng nếu mất kết nối/dữ liệu không kiểm soát hoặc tool lỗi; tiếp tục sau phục hồi và ghi log. Không đổi assertion để che lỗi.
## 6. Rủi ro
Cả hai ứng dụng do cùng AI hỗ trợ, không blind; công bố điều này. DB tạm cách ly để không phá dữ liệu; pilot bị lỗi không dùng làm benchmark cuối. Chưa xác định điểm bão hòa vì chỉ4users và ratecap.
## 7. Lịch trình
Ngày24–25/10: chuẩn hóa yêu cầu, dựng ứng dụng, test API/UI, load/step đúng thời lượng, đóng gói. Demo/chữ ký theo lịch hai sinh viên và giảng viên, chưa tự xác nhận.
## 8. Bàn giao
ReviewSRS, workbook7sheet, testcase/log/bugreport, Postman/environment/CSV/HTML, Puppeteer/data/ảnh, JMX/JTL/nativeHTML, summary và biên bản số2.
'''
    document(folder,f'Mau_04_Test_Plan_{msv}','KẾ HOẠCH KIỂM THỬ – MẪU 04',plan,person,msv,project,folder)
    summary=intro+'## 1. Tóm tắt\nĐã thực thi testcase và giữ nguyên FAIL khi phát hiện sai lệch. Số lỗi độc lập phân biệt với số testcase FAIL. Kết luận NO-GO nghiệm thu nghiệp vụ; đây là bản để thực hành kiểm thử.\n## 2. Kết quả thực thi\n'+writer.table(['Chỉ số','Chức năng','NFR'],[['Tổng',len(functional),nfr['total']],['Pass',fc['PASS'],nc.get('PASS',0)],['Fail',fc['FAIL'],nc.get('FAIL',0)],['Error/Blocked',fc['ERROR'],nc.get('ERROR',0)+nc.get('NOT_RUN',0)]])+'\n'
    summary+='### Độ phủ\n'+writer.table(['Mã FR','TC','Pass','Fail'],[(fr,len(rs:= [r for r in functional if r['function_id']==fr]),sum(r['status']=='PASS' for r in rs),sum(r['status']=='FAIL' for r in rs)) for fr in sorted({r['function_id'] for r in functional})])+'\n'
    summary+='## 3. Thống kê lỗi\n'+writer.table(['Loại','Major','Minor','Tổng'],[(kind,sum(b['category']==kind and b['severity']=='Major' for b in bugs),sum(b['category']==kind and b['severity']=='Minor' for b in bugs),sum(b['category']==kind for b in bugs)) for kind in ['Functional','Visual','Content','Performance','Crash']])+f'\nCó {len(bugs)} bug report; 20 lỗi cài sẵn và3phát hiện ngoài đặc tả. Không tính tỉ lệ phát hiện mù vì đã biết bản cài lỗi.\n'
    summary+='## 4. Kiểm thử tự động\n### API – Postman/Newman\n'+writer.table(['Chỉ số','Giá trị'],[(k,newman.get(k,'Xem raw summary')) for k in ['requests','iterations','assertions','assertionsPassed','assertionsFailed']])+'\n### Puppeteer\n'+writer.table(['Chỉ số','Giá trị'],[(k,pupp.get(k,'Xem raw summary')) for k in ['assertions','pass','fail','blocked']])+'\nCollection/CSV/scripts và ảnh/raw summary trong tooling/evidence. Đây là các lần thực chạy, không phải số dự kiến.\n'
    summary+='## 5. JMeter\nMôi trường: serverWerkzeug và máy tạo tải cùng Windows; SQLite tạm. Load2users300s; step1/2/4mỗi120s; timer180sample/phút. Không chạy stress hạ tầng trường.\n'+writer.table(['NFR','Kịch bản','Request','Ngưỡng P95ms','P95ms','Lỗi%','Kết luận'],[(r['requirement'],r['scenario'],r['sampler'],r['threshold_ms'],r['p95_ms'],r['error_percent'],r['result']) for r in perf])+'\n'
    for kind in ('load','step'):
        raw=read(toolroot/('jmeter-'+kind)/'summary.json',{})
        summary+=f'### Kịch bản {kind}\nThời lượng thực {raw.get("wall_seconds","chưa chốt")}s; đủ thời lượng: {raw.get("duration_verified",False)}.\n'
        for image,caption in [('response-times-over-time.png','Response Times Over Time'),('throughput-over-time.png','Throughput'),('active-threads-over-time.png','Active Threads')]:
            source=toolroot/('jmeter-'+kind)/image
            if source.exists():
                target=folder/'evidence/tooling'/('jmeter-'+kind)/image;copy(source,target)
                summary+=f'![Apache JMeter {kind}: {caption}](evidence/tooling/jmeter-{kind}/{image})\n'
    summary+='''Search/Notifications chậm ngay tải nhẹ, vượt ngưỡng800/900ms. Không suy ra bão hòa từ throughput3request/s vì đây là cap cấu hình; rawJTL là nguồn đếm người dùng, không suy sốVU từ đường nối đồ thị. Đề xuất bỏ độ trễ và đo lại sau demo; giữ nguyên bản lỗi cho bài tập.
## 6. Giới hạn
Đã biết mã nguồn/lỗi; không blind. Chưa đo Internet/VPS, chưa kiểm thử ngoài Edge, chưa stress trên4users; không cam kết hiệu năng production. Các chữ ký và sự phê duyệt của sinh viên/giảng viên chưa được thay thế.
## 7. Kết luận
NO-GO nghiệm thu nghiệp vụ do còn Major và các NFR không đạt. Hồ sơ có thể dùng để tái hiện và demo lỗi; không hợp thức hóa FAIL thành PASS.
## 8. Khai báo AI
Codex hỗ trợ dựng cả hai ứng dụng, thiết kế/viết/chạy công cụ, tổng hợp báo cáo. Tester là người phụ trách được phân công trong cặp, không phải tuyên bố đã chạy tay mọi testcase. Raw timestamp/hash/request/response/JTL/ảnh được giữ để kiểm chứng.
'''
    document(folder,f'Mau_05_Test_Summary_Report_{msv}','BÁO CÁO TỔNG KẾT KIỂM THỬ – MẪU 05',summary,person,msv,project,folder)
    hand=intro+'## BIÊN BẢN SỐ2 – BÀN GIAO KẾT QUẢ\nBên giao tester: '+person+'. Bên nhận chủ hệ thống: '+owner+'. Đây là hồ sơ chuẩn bị để hai bên rà soát, chưa ký nhận thay.\n## I. Danh mục\nReviewSRS,TestPlan,workbook/testlog/bugreports,Postman+HTML,Puppeteer+ảnh,JMX+JTL+HTML,Summary đều đi kèm trong thư mục này.\n## II. Kết quả\n'+f'Functional {len(functional)}: {fc["PASS"]}Pass,{fc["FAIL"]}Fail; NFR{nfr["total"]}: {nc.get("PASS",0)}Pass,{nc.get("FAIL",0)}Fail; bug{len(bugs)}. Kết luậnNo-go nghiệp vụ.\n## III. Ý kiến và chữ ký\nChưa ghi nhận ý kiến bất đồng/phê duyệt của hai sinh viên. BÊN GIAO: chưa ký. BÊN NHẬN: chưa ký.\n'
    document(folder,f'Mau_06_Bien_ban_ban_giao_2_{msv}','BIÊN BẢN BÀN GIAO KẾT QUẢ – MẪU 06',hand,person,msv,project,folder)

def main():
    normalize_tooling()
    shilo_functional=read(SUB/'partner/data/functional_cases.json')
    shilo_bugs=shilo_catalog();shilo_nf=shilo_nfr(shilo_functional)
    configs=[('Bui_Xuan_Minh_23110193','Bùi Xuân Minh','23110193','StudentCare',SUB/'public',SUB/'private/Mau_02_Bug_Manifest_23110193.xlsx',SUB/'evidence',read(SUB/'data/bug_catalog.json'), 'ShiloSkin','Nguyễn Từ Lân –23110306',shilo_functional,shilo_nf,shilo_bugs,SUB/'partner/evidence',PEER/'tooling',PEER/'evidence/tooling'),
             ('Nguyen_Tu_Lan_23110306','Nguyễn Từ Lân','23110306','ShiloSkin',PEER/'docs/public',None,SUB/'partner/evidence',shilo_bugs,'StudentCare','Bùi Xuân Minh –23110193',read(SUB/'pair-runs/lan-tests-studentcare/data/functional_cases.json'),read(SUB/'evidence/nonfunctional/results.json'),read(SUB/'data/bug_catalog.json'),SUB/'evidence',SUB/'tooling',SUB/'evidence/tooling')]
    for slug,name,msv,own,public,manifest,ownevidence,ownbugs,targetapp,owner,fc,nfr,bugs,evidence,tooling,toolproof in configs:
        home=PAIR/slug;part1=home/('01_Chu_he_thong_'+own);part2=home/('02_Tester_'+targetapp)
        copy(public,part1);copy(ownevidence,part1/'evidence')
        if manifest:copy(manifest,part1/manifest.name)
        else:workbook(part1/f'Mau_02_Bug_Manifest_{msv}.xlsx',name,owner,own,[],{},ownbugs,True)
        copy(evidence,part2/'evidence');copy(tooling,part2/'tooling');copy(toolproof,part2/'evidence/tooling')
        if msv=='23110306':copy(SUB/'pair-runs/lan-tests-studentcare/evidence/functional',part2/'evidence/functional')
        workbook(part2/f'Mau_03_Bo_bieu_mau_Kiem_thu_{msv}.xlsx',owner,name,targetapp,fc,nfr,bugs)
        tester_docs(part2,name,msv,owner,targetapp,fc,nfr,bugs,toolproof)
        checklist=f'''# Hồ sơ {name} – {msv}

- Chủ hệ thống: **{own}**; tester của **{targetapp}**.
- Part1: URD,SRS≥3role/16FR/3flows/10NFR, API, bàn giao1 chưa ký, Manifest20seed và bằng chứng.
- Part2: {len(fc)}testcase chức năng≥60,9E2E,≥4kỹ thuật; {nfr['total']}NFRcase≥10; workbook7sheet,review,log,{len(bugs)}bugreport,TestPlan,Summary,bàn giao2 chưa ký.
- Newman/Puppeteer/JMeter và raw evidence có trong Part2; không ghi PASS cho lỗi được tìm thấy.
- Bug Manifest chỉ gửi riêng giảng viên. Chưa gửi form vì chưa có URL form nộp; chữ ký/điểm dành cho người thật xác nhận.

{DISCLOSURE}

Ứng dụng StudentCare: https://sv23.bcse-vju.com / https://github.com/minhhal1180/qlychtsv

Ứng dụng ShiloSkin: https://sv21.bcse-vju.com / https://github.com/minhhal1180/shiloskin (mirror vì repo gốc Lân chỉ có quyền đọc).
'''
        (home/'README.md').write_text(checklist,encoding='utf-8')
        document(home,f'Checklist_truoc_khi_nop_{msv}','CHECKLIST HỒ SƠ TRƯỚC KHI NỘP',checklist,name,msv,own,home)
        score='Sinh viên: '+name+' –'+msv+'. Chủ hệ thống: '+own+'. Bạn cặp: '+owner+'. Ngày bảo vệ: chờ lịch.\n'+writer.table(['Mã','Tiêu chí','Tối đa','Điểm'],[(a,b,c,'Giảng viên điền') for a,b,c in [('A1','URD',1),('A2','SRS',1.5),('A3','API/sản phẩm',.5),('A4','Lỗi/manifest',1),('B1','Review/plan',.5),('B2','Testcase',1.5),('B3','Bugreport/tỉ lệpháthiện',1.5),('B4','API/UItools',1),('B5','JMeter',1.5)]])+'\nĐiểm cộng/trừ,điều chỉnh và chữ ký giảng viên: chưa xác nhận, không tự điền.\n'
        document(home,f'Phieu_cham_diem_{msv}','PHIẾU CHẤM ĐIỂM – CHỜ GIẢNG VIÊN',score,name,msv,own,home)
        fingerprint={str(p.relative_to(home)).replace('\\','/'):__import__('hashlib').sha256(p.read_bytes()).hexdigest() for p in home.rglob('*') if p.is_file() and p.name!='SHA256SUMS.json'}
        dump(home/'SHA256SUMS.json',fingerprint)
        destination=SUB/'packages'/f'Bo_bao_cao_{slug}.zip';destination.parent.mkdir(exist_ok=True)
        with zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as archive:
            for p in home.rglob('*'):
                if p.is_file():archive.write(p,p.relative_to(home))
        print(slug,'functional',len(fc),'nfr',nfr['total'],'bugs',len(bugs),'files',len(fingerprint),'zip',destination.stat().st_size)

if __name__=='__main__':main()
