"""Build reviewable student deliverables from actual test evidence and declared SRS.

The course templates are view-only: this recreates their observed structure,
not an exported copy. Never invent a peer review, signature, or test result.
"""
from __future__ import annotations
import ast
import csv
import hashlib
import html
import json
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
BASE = ROOT / 'submission'
PUBLIC = BASE / 'public'
PRIVATE = BASE / 'private'
SELF = PRIVATE / 'self-test'
for folder in (PUBLIC, PRIVATE, SELF):
    folder.mkdir(parents=True, exist_ok=True)
IDENTITY = 'Bùi Xuân Minh – 23110193'
DAY = '25/10/2026'
PARTNER_FILE = BASE / 'partner.json'
PARTNER = json.loads(PARTNER_FILE.read_text(encoding='utf-8')) if PARTNER_FILE.exists() else {}
PEER = PARTNER.get('name', 'Chờ thông tin bạn cặp – chưa xác nhận bàn giao')
REPO = 'https://github.com/minhhal1180/qlychtsv'
LIVE = 'https://sv23.bcse-vju.com'
PUBLIC_LINK = REPO + '/tree/main/submission/public'

FR = [
('FR-01','Đăng nhập','US-01','Cả 3 vai trò','Tài khoản 1–80 ký tự, mật khẩu 1–200; không được trống sau trim. Kiểm tra hash và active; sai/khóa/không tồn tại dùng cùng thông báo HTTP401. Đăng nhập đúng trả200, phiên đúng vai trò.'),
('FR-02','Phiên và đăng xuất','US-01','Cả 3 vai trò','Phiên ký số tồn tại tối đa8 giờ. /me trả thông tin và CSRF. Đăng xuất thu hồi các phiên tài khoản; cookie cũ gọi lại trả401; không xử lý nghiệp vụ sau hết phiên.'),
('FR-03','Quản lý tài khoản','US-09','ADMIN','Tạo username duy nhất1–80, name1–120, password8–200, role STUDENT/STAFF/ADMIN, active boolean. Admin sửa tên/role/active hoặc đặt mật khẩu mới; không đổi username; không trả hash qua API.'),
('FR-04','Vai trò và khóa','US-09','ADMIN','Backend kiểm tra vai trò trên mọi API. Khóa/đổi tài khoản thu hồi phiên. Admin không tự khóa hoặc bỏ quyền ADMIN. Cán bộ còn hồ sơ IN_PROGRESS/WAITING_STUDENT không đổi sang role khác.'),
('FR-05','Danh mục loại','US-10','ADMIN','Mã duy nhất1–30, tên1–120, mô tả0–1000, active boolean. Chỉ chọn loại active khi tạo mới, kể cả loại bị tắt sau khi mở form. Không xóa vật lý loại đã dùng; hồ sơ cũ vẫn tham chiếu.'),
('FR-06','Tạo yêu cầu','US-02','STUDENT','Loại active, tiêu đề1–200 sau trim, nội dung1–10000 sau trim. Tiêu đề toàn khoảng trắng phải400. Tạo201, đúng1 mã HT-xxxxx duy nhất, student hiện tại, NEW và thời gian; tệp tùy chọn. Retry cùng Idempotency-Key không tạo trùng.'),
('FR-07','Danh sách và tìm kiếm','US-03','STUDENT/STAFF','Sinh viên chỉ thấy của mình; cán bộ thấy hàng đợi chung, chọn mine=1 để lọc người phụ trách. Lọc6 trạng thái hợp lệ. Tìm chuỗi con theo mã HOẶC tiêu đề; Unicode casefold, không bỏ dấu. Số đếm hiện trạng trước lọc q/status.'),
('FR-08','Chi tiết và quyền sở hữu','US-03','STUDENT/STAFF','Chi tiết có loại, sinh viên, người phụ trách, trạng thái, ngày tạo/cập nhật, nội dung, trao đổi/tệp/lịch sử. Sinh viên không được đọc hồ sơ/tệp người khác dù biết ID:403; ADMIN không xử lý hồ sơ.'),
('FR-09','Hủy trước tiếp nhận','US-04','STUDENT','Chỉ chủ yêu cầu NEW được hủy. Chuyển CANCELLED nguyên tử và ghi lịch sử. Nếu cán bộ đã nhận trước, từ chối409 và giữ IN_PROGRESS. Hủy lần2 bị409.'),
('FR-10','Tiếp nhận','US-06','STAFF','Chỉ NEW chưa có assignee được nhận. Gán cán bộ+IN_PROGRESS trong cùng giao dịch. Hai cán bộ đồng thời: đúng1 thành công, người kia409. Báo cho sinh viên và ghi ACCEPT.'),
('FR-11','Yêu cầu bổ sung','US-07','STAFF được gán','Chỉ IN_PROGRESS, bắt buộc body1–10000 sau trim. Lưu lời yêu cầu và chuyển WAITING_STUDENT, ghi người/thời điểm/trước/sau; thông báo sinh viên.'),
('FR-12','Sinh viên bổ sung','US-05','STUDENT chủ hồ sơ','Chỉ WAITING_STUDENT. Có nội dung hoặc tệp hợp lệ. Lưu Message/Attachment và chuyển về IN_PROGRESS, ghi SUPPLEMENT, báo cán bộ; không sửa/xóa nội dung cũ.'),
('FR-13','Trao đổi và tệp','US-05','Chủ hồ sơ/cán bộ được gán','Cán bộ trao đổi trong IN_PROGRESS/WAITING_STUDENT; sinh viên bổ sung khi chờ. Tối đa3 tệp/lần,5MiB/tệp,16MiB request; PDF/PNG/JPG/TXT UTF-8, kiểm tra header; file rỗng400. Tải có xác thực và quyền từng bản ghi; tệp không tồn tại404.'),
('FR-14','Hoàn tất và từ chối','US-08','STAFF','COMPLETE chỉ IN_PROGRESS, bắt buộc kết quả1–10000, chuyển COMPLETED. REJECT từ NEW/IN_PROGRESS, bắt buộc lý do1–10000, chuyển REJECTED. Chỉ người phụ trách xử lý IN_PROGRESS; NEW được cán bộ tự từ chối. Mỗi chuyển trạng thái có lịch sử.'),
('FR-15','Mở lại','US-04','STUDENT chủ hồ sơ','Chỉ COMPLETED được mở lại trực tiếp. REJECTED phải409. Giữ cán bộ cũ nếu active và role STAFF; nếu không bỏ assignee, trả NEW. Có REOPEN; chưa giới hạn số lần/thời gian.'),
('FR-16','Thông báo nội bộ','US-11','Người nhận','Sau sự kiện thành công có thông báo đúng người nhận và liên kết hồ sơ. /read chỉ thông báo của mình; sau refresh is_read vẫn1. Retry thao tác không tạo thông báo lặp.'),
]
NFR = [
('NFR-01','Hiệu năng tìm kiếm','P95 ≤800ms cho GET /api/requests?q=... tại1–2 người dùng; lỗi<1%; tải thường300giây; dữ liệu10 yêu cầu (5seed +5tạo bởi Newman/Puppeteer trước đo). Đo response time HTTP, không cộng thời gian định thời tải.','JMeter/Python HTTP'),
('NFR-02','Hiệu năng thông báo','P95 ≤900ms cho GET /api/notifications tại1–2 người dùng; lỗi<1%; tải thường300giây.','JMeter/Python HTTP'),
('NFR-03','Hiệu năng danh sách','GET /api/requests không q và /api/types: P95 ≤1000ms, lỗi<1% ở các bậc1→2→4 người dùng, mỗi bậc120giây; ramp-up5giây/bậc, Constant Throughput180sample/phút cho cả nhóm (mục tiêu≤3request/giây); không suy rộng quá4 người dùng.','JMeter load+step'),
('NFR-04','Hiển thị','Viewport375×812,768×1024,1366×768: hộp thoại nằm trong viewport; nhãn trạng thái có thể nhìn thấy; tiêu đề hợp lệ có khoảng cách không bị cắt tại desktop. Bảng có thể cuộn trong vùng riêng.','Puppeteer/ảnh và bounding box'),
('NFR-05','Nội dung','100% nhãn6 trạng thái khớp từ điển: NEW=Mới gửi, IN_PROGRESS=Đang xử lý, WAITING_STUDENT=Chờ bổ sung, COMPLETED=Hoàn tất, CANCELLED=Đã hủy, REJECTED=Từ chối. Help tệp ghi5MB, không50MB. Ngày hiển thị dd/mm/yyyy theo vi-VN, giờ theo trình duyệt.','Đối chiếu DOM với dữ liệu'),
('NFR-06','Bảo mật cơ bản','0 lần thành công trong≥6 phép thử trái quyền/thiếu phiên/CSRF. HTTPS; cookie Secure,HttpOnly,SameSite=Lax; dữ liệu riêng chỉ người có quyền; không trả hash/token trong tài liệu bằng chứng.','API âm và live HTTPS'),
('NFR-07','Nhất quán và retry','Hai cán bộ đồng thời nhận1 NEW: đúng1 success; retry cùng key/body trả cùng id và chỉ1 CREATE/notification. Không mất giao dịch đã commit sau request lỗi khác.','ThreadPool/SQLite assertions'),
('NFR-08','Khả năng chịu dữ liệu lỗi','File0byte→400; file ID không tồn tại→404; 0 response500 cho hai trường hợp. Lỗi chỉ của request, không được dừng tiến trình; ngay sau lỗi /me và danh sách vẫn200.','API và HTTP'),
('NFR-09','Giới hạn tệp','Mỗi file1..5242880byte, tối đa3 tệp/lần; request≤16777216byte. Extension/header phù hợp; TXT UTF-8. Từ chối trường hợp sai không tạo hồ sơ/tệp một phần.','BVA và rollback assertions'),
('NFR-10','Truy vết và bất biến','100% chuyển trạng thái có actor,time,from,to; CREATE/ACCEPT/ASK/SUPPLEMENT/COMPLETE/REJECT/CANCEL/REOPEN. 0 API sửa/xóa message/history trong MVP.','API history, route inventory'),
]
US = [
('US-01','Cả3vai trò','đăng nhập đúng quyền và đăng xuất an toàn','dùng chức năng được cấp','P0','FR-01,FR-02','Thông tin đúng cho vào đúng vai; sai/quá phiên không xử lý.'),
('US-02','Sinh viên','gửi loại,tiêu đề,nội dung và tệp','nhờ nhà trường giải quyết một việc cụ thể','P0','FR-06','Nhận mã duy nhất NEW; lỗi nhập liệu được chỉ rõ.'),
('US-03','Sinh viên','tìm và xem tiến trình yêu cầu của mình','biết ai phụ trách và kết quả','P0','FR-07,FR-08','Tìm được bằng mã/tiêu đề; không xem dữ liệu sinh viên khác.'),
('US-04','Sinh viên','hủy khi chưa tiếp nhận hoặc mở lại hồ sơ hoàn tất','điều chỉnh nhu cầu hỗ trợ','P0','FR-09,FR-15','Hủy chỉNEW; mở lại chỉCOMPLETED và có lịch sử.'),
('US-05','Sinh viên','bổ sung lời giải thích/tệp theo đề nghị','cán bộ có đủ căn cứ','P0','FR-12,FR-13','Thông tin gửi bất biến, trở lại đang xử lý.'),
('US-06','Cán bộ','xem hàng đợi chung và tự nhận','đảm nhận một yêu cầu chưa xử lý','P0','FR-07,FR-10','Một hồ sơ chỉ có1cán bộ nhận thành công.'),
('US-07','Cán bộ','đề nghị sinh viên bổ sung khi thiếu','xử lý đúng dữ liệu','P0','FR-11','Bắt buộc giải thích; sinh viên thấy phản hồi.'),
('US-08','Cán bộ','ghi kết quả hoặc lý do từ chối','kết thúc minh bạch','P0','FR-14','Kết quả/lý do bắt buộc, có trạng thái và lịch sử.'),
('US-09','Quản trị viên','cấp/khóa tài khoản và đổi vai trò','quản lý người dùng hệ thống','P0','FR-03,FR-04','Chặn sai quyền; khóa thu hồi phiên cũ.'),
('US-10','Quản trị viên','quản lý danh mục và ngừng loại cũ','giữ danh mục phù hợp','P1','FR-05','Loại đã ngừng không tạo mới; hồ sơ cũ còn tham chiếu.'),
('US-11','Người nhận','đọc thông báo nội bộ','không bỏ lỡ cập nhật','P1','FR-16','Chỉ nhận của mình; trạng thái đã đọc được lưu.'),
]


def table(headers, rows):
    def cell(value):
        return str(value).replace('|','/').replace('\n','<br>')
    return '\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join('---' for _ in headers)+' |']+
                     ['| '+' | '.join(map(cell,row))+' |' for row in rows])


def write_document(stem, title, body, folder=PUBLIC):
    # Keep prose readable while artifact identifiers remain stable.
    for before, after in {
        'bạncặp':'bạn cặp','tựkiểmthử':'tự kiểm thử','tựkiểmtra':'tự kiểm tra',
        'chủhệthống':'chủ hệ thống','hệthống':'hệ thống','giảngviên':'giảng viên',
        'hồsơ':'hồ sơ','triểnkhai':'triển khai','hiệunăng':'hiệu năng','môitrường':'môi trường',
        'kếtquả':'kết quả','chưathựchiện':'chưa thực hiện','độclập':'độc lập',
        'pháthiện':'phát hiện','kiểmthử':'kiểm thử','kiểmtra':'kiểm tra','chứcnăng':'chức năng',
        'nghiệpvụ':'nghiệp vụ','dữliệu':'dữ liệu','tựđộng':'tự động','sảnphẩm':'sản phẩm',
        'kịchbản':'kịch bản','trướcdemo':'trước demo','chínhthức':'chính thức',
        'lỗicàisẵn':'lỗi cài sẵn','bằngchứng':'bằng chứng','mỗibậc':'mỗi bậc',
        'tỉlệ':'tỉ lệ','sốđo':'số đo','hạtầng':'hạ tầng','Tảithường':'Tải thường',
        'ngườidùng':'người dùng','độtrễ':'độ trễ','Độtrễ':'Độ trễ','cốý':'cố ý',
        'nhậnxét':'nhận xét','đánhgiá':'đánh giá','kếtluận':'kết luận','Kếtluận':'Kết luận',
        'biểuđồ':'biểu đồ','toànbộ':'toàn bộ','thựctế':'thực tế','thựchiện':'thực hiện',
        'thựcsự':'thực sự','đọclại':'đọc lại','Sinhviên':'Sinh viên','chữký':'chữ ký',
        'kýnhận':'ký nhận','đăngký':'đăng ký','bảovệ':'bảo vệ','côngcụ':'công cụ',
        'BùiXuânMinh':'Bùi Xuân Minh','hỗtrợ':'hỗ trợ','cóthể':'có thể','cụcbộ':'cục bộ',
        'đăngnhập':'đăng nhập','tiêuđề':'tiêu đề','chitiết':'chi tiết','thôngbáo':'thông báo',
        'danhsách':'danh sách','hàngđợi':'hàng đợi','thànhcông':'thành công',
        'khôngđạt':'không đạt','CHƯAHOÀNTẤT':'CHƯA HOÀN TẤT','khôngGo':'không Go',
        'SRS1.1':'SRS 1.1','URD1.0':'URD 1.0','TestPlan1.0':'Test Plan 1.0',
        'cóxácnhận':'có xác nhận','gửiriêng':'gửi riêng','chưathực':'chưa thực',
        'tài liệuPartA':'tài liệu Part A','tài khoảnvà':'tài khoản và',
        'hồ sơseed':'hồ sơ seed','bảnmanifestmới':'bản manifest mới',
        'cóhashfile':'có hash file','thuộcactive':'thuộc active'
    }.items():
        body = body.replace(before, after)
    full=f'# {title}\n\n{IDENTITY} · {DAY} · StudentCare v1.1\n\n'+body.strip()+'\n'
    path=folder/(stem+'.md'); path.write_text(full,encoding='utf-8')
    document=Document()
    document.sections[0].top_margin=Cm(1.8);document.sections[0].bottom_margin=Cm(1.8)
    normal=document.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    document.styles['Title'].font.color.rgb=RGBColor.from_string('176D53')
    lines=full.splitlines();index=0
    while index<len(lines):
        line=lines[index]
        if line.startswith('|') and index+1<len(lines) and re.match(r'^\|[\s:|\-]+\|$',lines[index+1]):
            data=[]
            while index<len(lines) and lines[index].startswith('|'):
                row=lines[index]
                if not re.match(r'^\|[\s:|\-]+\|$',row):data.append([c.strip().replace('<br>','\n') for c in row.strip('|').split('|')])
                index+=1
            grid=document.add_table(rows=1,cols=len(data[0]));grid.style='Light Shading Accent 1'
            for cell,text in zip(grid.rows[0].cells,data[0]):cell.text=text
            for row in data[1:]:
                for cell,text in zip(grid.add_row().cells,row):cell.text=text
            continue
        if line.startswith('# '):document.add_heading(line[2:],0)
        elif line.startswith('## '):document.add_heading(line[3:],1)
        elif line.startswith('### '):document.add_heading(line[4:],2)
        elif line.strip() and not line.startswith('```'):document.add_paragraph(line)
        index+=1
    footer=document.sections[0].footer.paragraphs[0];footer.text='StudentCare · 23110193 · Tài liệu chuẩn bị bằng AI, số liệu lấy từ bằng chứng thực thi'
    footer.runs[0].font.size=Pt(8)
    document.save(folder/(stem+'.docx'))
    pdf(full,folder/(stem+'.pdf'))


def pdf(text,path):
    if 'ArialUnicode' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('ArialUnicode','C:/Windows/Fonts/arial.ttf'))
        pdfmetrics.registerFont(TTFont('ArialUnicodeBold','C:/Windows/Fonts/arialbd.ttf'))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='VN',fontName='ArialUnicode',fontSize=9,leading=14,spaceAfter=6))
    styles.add(ParagraphStyle(name='VNTitle',fontName='ArialUnicodeBold',fontSize=21,leading=27,spaceAfter=18,textColor=colors.HexColor('#176d53')))
    styles.add(ParagraphStyle(name='VNHead',fontName='ArialUnicodeBold',fontSize=13,leading=18,spaceBefore=13,spaceAfter=8))
    story=[]
    for line in text.splitlines():
        if line.startswith('```') or re.match(r'^\|[\s:|\-]+\|$',line):continue
        if not line.strip():story.append(Spacer(1,4));continue
        style=styles['VNTitle'] if line.startswith('# ') else styles['VNHead'] if line.startswith('##') else styles['VN']
        content=html.escape(re.sub(r'^#+\s*','',line)).replace('&lt;br&gt;','<br/>')
        if line.startswith('|'):content=content.strip(' |').replace(' | ',' · ')
        story.append(Paragraph(content,style))
    def foot(canvas,doc):
        canvas.setFont('ArialUnicode',8);canvas.setFillColor(colors.grey)
        canvas.drawString(1.6*cm,1*cm,'StudentCare · Bùi Xuân Minh · 23110193')
        canvas.drawRightString(19.4*cm,1*cm,str(doc.page))
    SimpleDocTemplate(str(path),pagesize=(21*cm,29.7*cm),rightMargin=1.7*cm,leftMargin=1.7*cm,topMargin=1.6*cm,bottomMargin=1.6*cm).build(story,onFirstPage=foot,onLaterPages=foot)


def build_urd_srs():
    body='''## 1. Bối cảnh, mục tiêu và phạm vi
Sinh viên cần một điểm gửi câu hỏi, biết tiến trình và kết quả; một bộ phận hỗ trợ cần hàng đợi có người phụ trách và lịch sử. Mục tiêu: tạo → tiếp nhận → trao đổi → hoàn tất → theo dõi, không thay thế hệ thống đào tạo/tài chính.
Trong phạm vi: tài khoản riêng; một bộ phận; yêu cầu hỗ trợ; tệp; danh mục; thông báo nội bộ. Ngoài phạm vi: SSO/SIS, nhiều phòng ban, SLA động, email/SMS, chatbot, FAQ, dashboard chuyên sâu và xuất Excel.
Nguồn nghiệp vụ gốc: dactadan.pdf, Business Analysis Draft v1.0, 7 trang. Đây là URD được cụ thể hóa cho bản thực hành; các quyết định TBD là giả định được nêu trong SRS, không giả là chữ ký phê duyệt của nhà trường.
## 2. Vai trò và nhu cầu
'''+table(['Vai trò','Nhu cầu','Giới hạn'],[
('Sinh viên','Gửi, theo dõi, hủy NEW, bổ sung, mở lại COMPLETED','Chỉ dữ liệu của mình'),('Cán bộ hỗ trợ','Hàng đợi chung, nhận, hỏi bổ sung, kết thúc','Xử lý hồ sơ đã được gán'),('Quản trị viên','Tài khoản, vai trò, danh mục','Không dùng API xử lý yêu cầu')])+'''
## 3. User story có độ ưu tiên
'''+table(['ID','User story','Ưu tiên','Chức năng','Nghiệm thu'],[(i,f'Là {actor}, tôi muốn {want}, để {why}.',priority,fr,ac) for i,actor,want,why,priority,fr,ac in US])+'''
## 4. Luồng nghiệp vụ từ góc nhìn người dùng
FLOW-01 – Giải quyết thông thường: sinh viên gửi hồ sơ hợp lệ và nhận mã; cán bộ nhận hồ sơ; cán bộ ghi kết quả; sinh viên thấy hoàn tất và thông báo. Thay thế: sinh viên hủy trước tiếp nhận; ngoại lệ: hồ sơ đã có người nhận, không được hủy.
FLOW-02 – Bổ sung và tiếp tục: cán bộ nhận, chỉ rõ thiếu gì; sinh viên gửi nội dung/tệp; trạng thái quay về đang xử lý; cán bộ ghi kết quả. Ngoại lệ: sinh viên khác hoặc tệp sai bị chặn, lịch sử cũ giữ nguyên.
FLOW-03 – Từ chối và mở lại: cán bộ nêu lý do từ chối hồ sơ không thuộc phạm vi; hồ sơ từ chối không mở lại. Với hồ sơ đã hoàn tất, chủ hồ sơ đề nghị mở lại trực tiếp; giữ cán bộ còn hợp lệ hoặc quay về hàng đợi.
## 5. Tiêu chí nghiệm thu và ràng buộc
Đủ3vai trò,16chức năng,3luồng; người dùng chỉ dùng quyền được cấp; mọi kết thúc có kết quả hoặc lý do và lịch sử. Mức đo hiệu năng, file, phiên và giao diện được quy định cụ thể ở SRS. Bản thực hành được chấp nhận để học kiểm thử; không kết luận đạt nghiệp vụ khi còn lỗi.
## 6. Ghi nhận và phê duyệt
Người lập: Bùi Xuân Minh – 23110193. Hỗ trợ AI: Codex tạo mã và bản thảo, sinh viên cần đọc, chạy lại và chịu trách nhiệm nội dung. Bạn cặp: '''+PEER+'''. Tình trạng: bản chuẩn bị bàn giao; chưa ghi nhận chữ ký hoặc phê duyệt của bạn cặp/giảng viên.
'''
    write_document('URD_23110193','TÀI LIỆU YÊU CẦU NGƯỜI DÙNG (URD)',body)
    srs='''## 1. Kiểm soát tài liệu và căn cứ
Phiên bản1.1 – cụ thể hóa Draft gốc cho đồ án. Baseline yêu cầu đúng, không mô tả hành vi lỗi là đúng. URD v1.0 ngày24/10/2026; dactadan.pdf; checklist và mẫu Drive. Chưa có review độc lập bạn cặp. Các mục bổ sung giao diện/hiệu năng là yêu cầu bài thực hành mới, không gán ngược vào PDF gốc.
## 2. Actor và phân quyền
'''+table(['Thao tác','STUDENT','STAFF','ADMIN'],[
('Tạo yêu cầu','Có','Không','Không'),('Xem chi tiết/tệp','Chỉ của mình','Hàng đợi chung','Không'),('Nhận NEW','Không','Có','Không'),('Xử lý IN_PROGRESS','Không','Chỉ được gán','Không'),('Bổ sung WAITING_STUDENT','Chủ hồ sơ','Trao đổi nếu được gán','Không'),('Hủy/mở lại','Đúng trạng thái,của mình','Không','Không'),('Tài khoản/danh mục','Không','Không','Có'),('Thông báo','Chỉ người nhận','Chỉ người nhận','Chỉ người nhận')])+'''
## 3. Yêu cầu chức năng và truy vết
'''
    for fid,name,usid,actors,rule in FR:
        srs+=f'### {fid} – {name}\nTruy vết {usid}. Actor: {actors}.\n{rule}\nNghiệm thu: gọi giao diện/API theo quyền và dữ liệu trên, so sánh trạng thái DB/lịch sử; thao tác sai không được ghi một phần.\n'
    srs+='''## 4. Luồng hoàn chỉnh
### FLOW-01 – Tạo, nhận, hoàn tất
Actor chính: sinh viên/cán bộ. Tiền điều kiện: tài khoản active và loại active. Trigger: sinh viên cần xác nhận giấy tờ. Luồng chính: (1) đăng nhập, (2) gửi loại+tiêu đề+nội dung, (3) trả mã NEW, (4) cán bộ đọc hàng đợi và nhận, (5) IN_PROGRESS có assignee, (6) nhập kết quả, (7) COMPLETED và thông báo/lịch sử, (8) sinh viên xem kết quả. Hậu điều kiện:1hồ sơ,1người phụ trách, có CREATE/ACCEPT/COMPLETE.
Thay thế A1: trước bước4, sinh viên hủy → CANCELLED và CANCEL. Ngoại lệ E1: hai cán bộ nhận đồng thời chỉ1thành công. E2: kết quả rỗng→400. E3: timeout sau commit, retry cùngkey trả kết quả cũ.
### FLOW-02 – Chờ sinh viên bổ sung
Tiền điều kiện: hồ sơ IN_PROGRESS thuộc cán bộ. Luồng chính: (1) cán bộ ghi nội dung cần bổ sung, (2) WAITING_STUDENT, (3) sinh viên nhận thông báo, (4) gửi body hoặc tệp, (5) lưu bất biến, IN_PROGRESS và báo cán bộ, (6) cán bộ hoàn tất với kết quả. Hậu điều kiện: giữ trao đổi và đủ ASK/SUPPLEMENT/COMPLETE.
Thay thế A1: chỉ gửi tệp hợp lệ, body rỗng vẫn được. E1: cả nội dung và file rỗng→400. E2: sinh viên khác→403. E3: quá giới hạn/định dạng→400, không lưu một phần.
### FLOW-03 – Từ chối, hoàn tất và mở lại
Tiền điều kiện: hồ sơ NEW hoặc IN_PROGRESS. Luồng chính A: cán bộ nhập lý do từ chối → REJECTED, sinh viên xem; mở lại REJECTED→409. Luồng chính B: hồ sơ đã COMPLETED, sinh viên mở lại → IN_PROGRESS với cán bộ cũ; ghi REOPEN. Thay thế B1: cán bộ cũ bị khóa/mất STAFF→NEW không assignee. E1: người khác mở lại→403. E2: mở lại NEW/CANCELLED/IN_PROGRESS→409. Không giới hạn số lần trong MVP.
## 5. Bảng quyết định và trạng thái
'''+table(['Hành động','Actor/điều kiện','Từ','Sang','Không hợp lệ'],[
('cancel','Chủ hồ sơ','NEW','CANCELLED','409 trạng thái/403 quyền'),('accept','STAFF,chưa gán','NEW','IN_PROGRESS','409'),('ask','STAFF được gán,body','IN_PROGRESS','WAITING_STUDENT','400/403/409'),('supplement','Chủ hồ sơ,body hoặc file','WAITING_STUDENT','IN_PROGRESS','400/403/409'),('complete','Được gán,kết quả','IN_PROGRESS','COMPLETED','400/403/409'),('reject','STAFF,lý do','NEW/IN_PROGRESS','REJECTED','400/403/409'),('reopen','Chủ hồ sơ','COMPLETED','IN_PROGRESS/NEW','403/409')])+'''
Sơ đồ trạng thái: NEW → IN_PROGRESS → WAITING_STUDENT → IN_PROGRESS → COMPLETED → IN_PROGRESS hoặc NEW. Nhánh kết thúc: NEW→CANCELLED; NEW/IN_PROGRESS→REJECTED. Trạng thái CANCELLED/REJECTED không nhận trao đổi mới. Điều kiện cập nhật và gán người được bảo vệ bởi SQLite BEGIN IMMEDIATE.
## 6. Phi chức năng đo được
Các ngưỡng dưới đây là MỤC TIÊU cần nghiệm thu, không phải kết quả đã đạt. Thiết lập tải nêu rõ ở mỗi mục; kết quả đo lấy từ JTL/JSON thật. Không đăng ký chạy tải trên hạ tầng trường nếu chưa được cấp khung giờ; bộ số đo chuẩn bị chạy trên bản sao cục bộ và không đại diện VPS.
'''+table(['Mã','Nhóm','Tiêu chí có số đo','Phương pháp'],NFR)+'''
## 7. Dữ liệu và validation
User(id,username,name,password_hash,role,active,session_version). RequestType(id,code,name,description,active). SupportRequest(id,code,student_id,type_id,title,content,status,assignee_id,created_at,updated_at). History, Message, Attachment, Notification theo PDF. Idempotency(user,method,path,key,digest,response,status) là bổ sung kỹ thuật.
Ràng buộc ký tự theo FR; ID là nguyên dương; boolean không nhận chuỗi. Tệp PDF bắt đầu %PDF-, PNG đúng signature, JPG FFD8FF, TXT giải mã UTF-8. Thời gian lưu ISO8601 UTC; UI trình bày vi-VN. Mã HT-xxxxx có thể tăng quá5chữ số khi đủ dữ liệu; không giả là độ dài cố định. Không phân trang trong MVP: cần xem lại khi dữ liệu lớn.
## 8. Quyết định TBD, giới hạn và review
Chọn hàng đợi chung; mở lại trực tiếp; giữ cán bộ cũ còn hợp lệ; admin cấp/đặt lại mật khẩu; file theo ngưỡng trên; chưa mở lại REJECTED; số đếm theo dữ liệu được phép xem trước q/status. Các đề xuất gốc được ghi thành lựa chọn bản thực hành, chưa là quyết định chính thức của nhà trường.
Thông báo và nghiệp vụ hiện cùng transaction; khác đề xuất PDF tách lỗi thông báo, chưa có outbox. Chưa có rate limit, backup tự động, monitoring tập trung. Hoạt động thực tế cần đánh giá thêm các giới hạn này. Review nội bộ đã cụ thể hóa trường, trạng thái, quyền, hiệu năng; chưa nhận review và xác nhận từ bạn cặp.
## 9. Truy vết
'''+table(['FR','URD','PDF gốc'],[(fr,us,'F01' if fr in ('FR-01','FR-02','FR-03','FR-04') else 'F05' if fr in ('FR-05','FR-16') else 'F04' if fr in ('FR-12','FR-13') else 'F03' if fr in ('FR-10','FR-11','FR-14') else 'F02') for fr,_,us,_,_ in FR])+'''
16chức năng,3luồng,10NFR; NFR-01..03 là hiệu năng. Testcase truy vết theo mã trong workbook; các bug khám phá ngoài đặc tả gốc được ghi riêng, không giả là yêu cầu đã phê duyệt.
'''
    write_document('SRS_23110193','ĐẶC TẢ YÊU CẦU PHẦN MỀM (SRS)',srs)


def endpoints():
    tree=ast.parse((ROOT/'app.py').read_text(encoding='utf-8'))
    found=[]
    for node in ast.walk(tree):
        if isinstance(node,ast.FunctionDef):
            for decorator in node.decorator_list:
                if isinstance(decorator,ast.Call) and isinstance(decorator.func,ast.Attribute) and decorator.func.attr in ('get','post','patch') and decorator.args and isinstance(decorator.args[0],ast.Constant):
                    path=decorator.args[0].value
                    if path.startswith('/api/'):
                        found.append((decorator.func.attr.upper(),path,node.name))
    return found


API_META = {
'login':('FR-01','Công khai',{'username':'sv01','password':'Demo@123'},{'ok':True},200,'username:string1..80 bắt buộc; password:string1..200 bắt buộc; không trống; phải khớp tài khoản active.'),
'logout':('FR-02','Đã đăng nhập',{}, {'ok':True},200,'Cookie phiên và X-CSRF-Token hợp lệ; thu hồi toàn bộ phiên của tài khoản.'),
'me':('FR-02','Đã đăng nhập',{}, {'user':{'id':1,'username':'sv01','name':'Nguyễn Minh Anh','role':'STUDENT'},'csrf':'<csrf>'},200,'Không body; session cookie. Không trả password_hash.'),
'types':('FR-05','Đã đăng nhập',{}, {'items':[{'id':1,'code':'HOCVU','name':'Học vụ','description':'Học vụ','active':1}]},200,'Admin thấy cả active/inactive; vai trò khác chỉ active.'),
'create_type':('FR-05','ADMIN',{'code':'KHAC','name':'Hỗ trợ khác','description':'Nội dung khác','active':True},{'id':4},201,'code:string1..30 duy nhất; name:string1..120; description:string0..1000; active:boolean tùy chọn mặc định true.'),
'edit_type':('FR-05','ADMIN',{'name':'Học vụ cập nhật','active':False},{'ok':True},200,'type_id:path nguyên dương tồn tại; các trường tùy chọn cùng giới hạn create.'),
'users':('FR-03','ADMIN',{}, {'items':[{'id':1,'username':'sv01','name':'Nguyễn Minh Anh','role':'STUDENT','active':1}]},200,'Không body; không trả hash.'),
'create_user':('FR-03','ADMIN',{'username':'sv03','name':'Sinh viên thử','password':'Password@123','role':'STUDENT','active':True},{'id':6},201,'username:string1..80 duy nhất; name:string1..120; password:string8..200; role enum; active boolean default true.'),
'edit_user':('FR-04','ADMIN',{'name':'Sinh viên cập nhật','active':False},{'ok':True},200,'user_id:path nguyên dương; name,role,active,password tùy chọn cùng validation. Không tự khóa/bỏ ADMIN; cán bộ có hồ sơ mở không đổi role.'),
'list_requests':('FR-07','STUDENT/STAFF',{}, {'items':[], 'counts':{s:0 for s in ['NEW','IN_PROGRESS','WAITING_STUDENT','COMPLETED','CANCELLED','REJECTED']}},200,'Query q:string0..200; status:1trong6enum hoặc rỗng; mine=1 lọc assignee với STAFF. Không phân trang.'),
'create_request':('FR-06','STUDENT',{'type_id':1,'title':'Xin xác nhận sinh viên','content':'Em cần hồ sơ học bổng.'},{'id':6,'code':'HT-00006'},201,'type_id:integer hoặc chuỗi ID nguyên dương active; title:string1..200 trim; content:string1..10000 trim. Multipart tùy chọn files≤3,≤5MiB/tệp.'),
'detail':('FR-08','STUDENT chủ hồ sơ/STAFF',{}, {'item':{'id':1,'code':'HT-00001','status':'NEW','title':'Xin xác nhận'},'messages':[],'history':[{'action':'CREATE','from_status':None,'to_status':'NEW','actor_name':'Nguyễn Minh Anh'}],'attachments':[]},200,'ticket_id:path nguyên dương. JSON mẫu rút gọn; item có student_id/type_id/assignee_id,ngày,content; messages có sender,time; history có actor/time/from/to.'),
'action':('FR-09/10/11/14/15','Theo bảng quyết định SRS',{'action':'accept','body':''},{'ok':True},200,'ticket_id:path. action enum accept/cancel/ask/complete/reject/reopen; body:string0..10000, bắt buộc có nội dung khi ask/complete/reject. Quyền và trạng thái theo SRS.'),
'message':('FR-12/13','Chủ hồ sơ/STAFF được gán',{'body':'Em đã gửi bổ sung.'},{'id':3},201,'ticket_id:path; body:string0..10000; body hoặc files bắt buộc. Multipart files≤3≤5MiB/tệp; check trạng thái theo SRS.'),
'download':('FR-13','Chủ hồ sơ/STAFF',{},'Binary; Content-Disposition: attachment; filename=<tên tệp>',200,'attachment_id:path nguyên dương tồn tại. Đây là endpoint nhị phân, không ép JSON cho response download.'),
'notifications':('FR-16','Người nhận đăng nhập',{}, {'items':[{'id':1,'request_id':1,'event_type':'CREATE','is_read':0,'code':'HT-00001','title':'Xin xác nhận'}]},200,'Không body; chỉ dữ liệu recipient hiện tại; sort ID giảm dần.'),
'read_notification':('FR-16','Người nhận',{}, {'ok':True},200,'notification_id:path; phải thuộc recipient, không thì404. Lưu is_read=1.'),
}


def build_api():
    api='''## 1. Thông tin chung
Base URL: https://sv23.bcse-vju.com/api (local http://127.0.0.1:5000/api). JSON UTF-8, Content-Type: application/json; upload multipart/form-data; download binary. Cookie session ký số; POST /login đặt cookie, GET /me lấy CSRF; POST/PATCH còn lại gửi X-CSRF-Token. Phiên8giờ; Secure/HttpOnly/SameSite=Lax trên HTTPS.
Thời gian ISO8601 UTC trong API; UI theo vi-VN. Không phân trang; không có Swagger. Không dùng Bearer. Ví dụ dưới là mẫu hợp đồng đúng; id/thời điểm minh họa, không phải số đo thực thi. Khác biệt actual được ghi ở báo cáo riêng, không hợp thức hóa lỗi bằng tài liệu.
## 2. Mã lỗi chung
'''+table(['HTTP','Ý nghĩa','Body mẫu'],[(200,'Thành công','{"ok":true} hoặc {"items":[]}'),(201,'Tạo mới','{"id":6,"code":"HT-00006"}'),(400,'Validation','{"error":"Thông báo validation tiếng Việt"}'),(401,'Chưa đăng nhập/hết phiên/sai thông tin','{"error":"Vui lòng đăng nhập."}'),(403,'Sai quyền/CSRF','{"error":"Bạn không có quyền thực hiện thao tác này."}'),(404,'Không tìm thấy','{"error":"Không tìm thấy tệp."}'),(409,'Trùng/xung đột trạng thái','{"error":"Yêu cầu đã được nhận hoặc đã kết thúc."}'),(413,'Body quá lớn','{"error":"The data value transmitted exceeds the capacity limit."}'),(500,'Lỗi máy chủ ngoài hợp đồng','Có thể là JSON hoặc HTML tùy nguồn; luôn xem status/content-type, không coi500 là lỗi nhập liệu hợp lệ.')])+'''
## 3. Danh sách đầy đủ endpoint
'''+table(['STT','Method','Endpoint','Tên','FR','Vai trò'],[(i,method,path,name,API_META[name][0],API_META[name][1]) for i,(method,path,name) in enumerate(endpoints(),1)])+'\n## 4. Đặc tả từng endpoint\n'
    for i,(method,path,name) in enumerate(endpoints(),1):
        fr,role,req,res,status,rule=API_META[name]
        api+=f'### 4.{i}. {method} {path}\nMethod + URL: {method} '+LIVE+path+f'\nChức năng: {fr}; vai trò: {role}.\nHeaders: Cookie phiên; '+('Content-Type theo JSON/multipart, X-CSRF-Token trừ /login; Idempotency-Key tùy chọn≤100ký tự.' if method!='GET' else 'Không body; GET dùng cookie.')+'\nRàng buộc path/query/body: '+rule+'\nRequest mẫu: '+json.dumps(req,ensure_ascii=False)+'\nResponse thành công '+str(status)+': '+json.dumps(res,ensure_ascii=False)+'\nResponse lỗi: 400 nếu validation (với endpoint có input);401 thiếu/hết phiên trừ login kiểm tra thông tin;403 sai vai trò/quyền/CSRF;404 ID không tồn tại;409 trùng/xung đột đối với thao tác ghi. Body {"error":"Thông báo cụ thể tiếng Việt"}. Với endpoint không phát sinh một mã, mã đó không phải kỳ vọng cần tạo.\n'
        if name=='action':
            api+='Request mẫu bổ sung từng action: '+json.dumps([{'action':a,'body':b} for a,b in [('cancel',''),('ask','Bổ sung bảng điểm'),('complete','Đã cấp giấy'),('reject','Không thuộc phạm vi'),('reopen','')]],ensure_ascii=False)+'\n'
    api+='''## 5. Tài khoản và dữ liệu mẫu
'''+table(['Vai trò','Tài khoản','Mật khẩu mẫu','Ghi chú'],[('STUDENT','sv01/sv02','Demo@123','Có dữ liệu riêng'),('STAFF','cb01/cb02','Demo@123','Cán bộ1 có3hồ sơ mẫu'),('ADMIN','admin','Demo@123','Quản lý tài khoản/danh mục')])+'''
Dữ liệu ban đầu5yêu cầu,3loại,5tài khoản. Retry dùng cùng Idempotency-Key, cùng path/body và tài khoản; multipart fingerprint không phụ thuộc boundary. Tải lại cùngkey không thêm hồ sơ/lịch sử/thông báo. Dữ liệu demo có thể tăng sau khi tester thao tác; khôi phục trên DB sao chép hoặc backup, không reset dữ liệu người khác.
'''
    write_document('Mau_01_Tai_lieu_API_23110193','TÀI LIỆU API – MẪU 01',api)


def load_json(path,default=None):
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else default


def build_plan_and_handover():
    plan='''## Trạng thái và phạm vi của bản này
Đây là KẾ HOẠCH TỰ KIỂM THỬ StudentCare của Bùi Xuân Minh, không phải xác nhận đã kiểm thử hệ thống bạn cặp. Khi có bàn giao độc lập, thay thông tin hệ thống, SRS, tài khoản và chạy lại toàn bộ công cụ; không sao chép kết quả tự kiểm thử thành kết quả tester.
Tester tự kiểm tra: '''+IDENTITY+'''. Chủ hệ thống: '''+IDENTITY+'''. Bạn cặp: '''+PEER+'''. Link: '''+LIVE+'''.
## 1. Giới thiệu
1.1. Mục đích: kiểm chứng chức năng/quyền/trạng thái và NFR của SRS1.1; phát hiện sai lệch, không sửa lỗi trước demo. Đánh giá khả năng bàn giao để bạn cặp kiểm thử; không kết luận nghiệm thu nghiệp vụ khi còn lỗi.
1.2. Tài liệu tham chiếu: URD1.0, SRS1.1, API1.1, dactadan.pdf, phiếu review nội bộ trong workbook. Chưa có review/chữ ký độc lập của bạn cặp.
## 2. Phạm vi
2.1. Trong phạm vi: toàn bộ FR-01..16; FLOW-01..03; NFR-01..10. Login, dữ liệu, quyền, rollback, file, phiên, retry, đồng thời; giao diện và nội dung3độ rộng; API tự động, UI tự động, JMeter tải thường/tăng tải.
'''+table(['Mã','Mức ưu tiên','Loại kiểm thử'],[(fr,'High','API + EP/BVA/Decision/State; E2E ở luồng tương ứng') for fr,*_ in FR]+[(nfr,'High' if nfr in ('NFR-01','NFR-02','NFR-03','NFR-06','NFR-07','NFR-08') else 'Medium',method) for nfr,_,_,method in NFR])+'''
2.2. Ngoài phạm vi: hệ thống bạn cặp khi chưa có thông tin; tải trên hạ tầng trường khi chưa có khung giờ; email/SMS/SSO/SIS/thanh toán thật/nhiều phòng ban; hiệu năng ngoài4người dùng; trình duyệt ngoài Edge/Chromium; đánh giá khả năng vận hành nghiệp vụ thật.
## 3. Chiến lược
3.1. EP: hợp lệ/trống/sai kiểu; BVA: giới hạn tiêu đề200,nội dung10000,file5MiB; bảng quyết định: quyền×hành động×trạng thái; chuyển trạng thái: các trạng thái kết thúc; Use case/E2E≥9case phủ3luồng; phân quyền trực tiếp API; exploratory thêm3heuristic UX ngoài PDF gốc.
API: Postman/Newman có≥6chức năng, data-driven3dòng, kiểm thử âm, assertionstatus/body. UI: Puppeteer dữ liệu riêng,≥2luồng, assertion+ảnh. Hiệu năng: JMeter gốc5.6.3, non-GUI, load2users300s/ramp10s; step1→2→4,120s/bậc,ramp5s/bậc; timer180sample/phút cho nhóm. Chạy read-only trên DB tạm, không ghi vào sản phẩm của người khác.
3.2. Loại lỗi: Functional – nghiệp vụ sai; Visual – bố cục/hiển thị; Content – nội dung/định dạng sai; Performance – vượt ngưỡng; Crash – HTTP500/trang lỗi/mất phiên đột ngột. Mức Critical: chặn luồng/mất dữ liệu/bảo mật/sập; Major: quan trọng sai, còn vòng tránh; Minor: phụ/giao diện; Trivial: thẩm mỹ. Priority độc lập với severity.
## 4. Môi trường
Server sản phẩm: Ubuntu24.04,Python3.12,Gunicorn26.2,Nginx,SQLite; linkHTTPS nêu trên. Máy tự kiểm thử/tạo tải: Windows,Python3.11.9,CPU16logical,serverWerkzeug threaded cùng máy; dữ liệu tạm10hồ sơ trước JMeter (5seed+5tạo bởi công cụ). Số đo này KHÔNG đại diện VPS.
Trình duyệt Edge155; viewport375/390,768,1366/1440. Node24.11.1,Newman6.2.1,htmlextra1.23.1,Puppeteer-core24.38.0,JMeter5.6.3,JDK21. Phiên bản thực lấy summaryJSON/log. Tài khoản mẫu:sv01/sv02,cb01/cb02,admin; Demo@123. DB tạm mỗi test độc lập, không reset DB live. Khung giờ JMeter: tự kiểm thử cục bộ bắt đầu24/10/2026,qua25/10 theo UTC+7; chưa đăng ký tải ở trường.
## 5. Tiêu chí
Bắt đầu: bản yêu cầu1.1 có thể viết testcase, tạo được DB tạm, mọi vai trò đăng nhập, phiên bản/hash được ghi. Kết thúc kỹ thuật:100%case có kết quả hoặc lý doBlocked/NotRun; có rawproof, báo cáo và liên kết; đã phân loại lỗi. Nghiệm thu nghiệp vụ Go chỉ khi0Critical/Major chưa xử lý và cácNFRđạt. Bản thực hành giữ lỗi nên kết luậnNo-go nghiệp vụ.
Tạm dừng: serverkhôngtruycập, dữ liệu bị đổi không thể kiểm soát, lỗi môi trường tool, vượt giới hạn tải cho phép. Tiếp tục: xác nhận môi trường phục hồi, gán lại dataset/hash, chạy lại case bị chặn; không thay kết quả cũ.
## 6. Rủi ro
'''+table(['Rủi ro','Ảnh hưởng','Giảm thiểu'],[('Chưa có tài liệu bạn cặp','Phần tester chính thức bị chặn','Giữ riêng self-test, chờ bàn giao thực'),('Phiên SSH hết hạn','Không cập nhật live','Xin slot mới, không đoán credentials'),('Dữ liệu dùng chung','Không lặp lại','DBtạm percase/backup live'),('Công cụ/servercùngmáy','Sai suy rộng hiệu năng','Ghi rõ môitrường, khônggọi sốđoVPS'),('Manifest lộ','Mất tính độc lập','Tệp riêng/private; không gửi cùng tài liệu bạn cặp'),('Bản10lỗicũđãcôngkhai','Không bảo đảm bí mật lịch sử','Không khẳng định blindtest; báo hạn chế cho giảng viên')])+'''
## 7. Lịch trình
'''+table(['Hoạt động','Thời gian/thực trạng','Sản phẩm'],[('Đọc yêu cầu/mẫu','24/10; đã đọc9mẫu qua viewer','Danh mụcnguồn'),('Chuẩn hóaURD/SRS/API','24–25/10','Bộ tài liệuPartA'),('Thiết kế/chạy testnội bộ','24–25/10','97functional+≥10NFR'),('API/UItựđộng','24–25/10','Newman/Puppeteerproof'),('Hiệunăngcụcbộ','24–25/10,đúngthờilượng','JTL/HTML'),('Bàn giao thật bạn cặp','Chờ tài liệu/thông tin','Biên bản1cóxácnhận'),('Tester hệ thốngbạn','Sau khi nhận đủ','BộPartBthật,không dùngselftest'),('Demo','Theo lịchgiảngviên,chưa có','Giữ nguyênlỗivàlink')])+'''
## 8. Sản phẩm bàn giao
ReviewSRS,TestPlan,workbook7sheets,rawfunctional/NFR,bugreport cóảnh,Postmancollection/environment/data/NewmanHTML,Puppeteerscript/data/ảnh,JMX/JTL/nativeHTMLdashboard,TestSummary. Dữ liệu thật lấy từ file, không tự điền tỉ lệ tìm lỗi độc lập khi chưa cóManifest của bạn cặp.
'''
    write_document('Mau_04_Test_Plan_TU_KIEM_THU_23110193','KẾ HOẠCH KIỂM THỬ – MẪU 04 (TỰ KIỂM THỬ)',plan,SELF)
    hand='''## BIÊN BẢN BÀN GIAO SỐ 1 – SẢN PHẨM
Bên giao: '''+IDENTITY+'''. Bên nhận: '''+PEER+'''. Hệ thống: StudentCare. Ngày lập dự thảo: '''+DAY+'''. Đây là bản sẵn sàng để hai bên kiểm tra, chưa phải xác nhận nhận đủ/chữ ký.
### I. Danh mục
'''+table(['STT','Hạng mục','Phiên bản/link','Ghi chú'],[(1,'URD','URD_23110193.docx / .pdf','v1.0'),(2,'SRS','SRS_23110193.docx / .pdf','v1.1; chờ reviewđộclập'),(3,'API','Mau_01_Tai_lieu_API_23110193.docx / .pdf','Tất cảendpoint'),(4,'Sản phẩm',LIVE,'Tài khoảnvàdữliệumẫu'),(5,'Tài khoản','sv01/sv02,cb01/cb02,admin; Demo@123','Mỗi3vai trò'),(6,'Dữ liệu','5tài khoản3loại5hồ sơseed','Sốlivecóthểtăngtheotest')])+'''
### II. Xác nhận bên nhận
Chưa đánh dấu các ô “Truy cập được link”, “Đăng nhập mọi vai”, “Nhận đủ tài liệu”: bên nhận tự kiểm tra và xác nhận. Kiểm tra tự động của chủ hệ thống không thay chữ ký bên nhận.
### III. Hạn chế đã biết
Một bộ phận; khôngSSO/email/SMS/dashboardchuyênsâu; chưabackupauto; khônghỗtrợvận hành thật. Tài liệu này không liệt kê lỗi cài sẵn. NFR là mục tiêu chưa được chứng nhận độc lập. Bản triển khai và cập nhật xem README; dữ liệu và khóa phiên giữ trên VPS, không ởGitHub.
### IV. Cam kết và ký nhận
Chủ hệ thống dự kiến giữ nguyên phiên bản bàn giao đến demo, duy trì link và phản hồi bugreport trong48giờ sau khi nhận. Lịch chính thức tuần6–14 do giảngviên xác nhận, không tự suy ngày từ hôm nay. BugManifest gửi riênggiảngviên, không gửi bạn test. Trước khi ký, hai bên đối chiếu tag/hash triển khai.
BÊN GIAO: Bùi Xuân Minh –23110193 (chưa ký). BÊN NHẬN: '''+PEER+''' (chưa ký).
'''
    write_document('Mau_06_Bien_ban_ban_giao_1_23110193','BIÊN BẢN BÀN GIAO SẢN PHẨM – MẪU 06',hand)
    hand2='''## BIÊN BẢN BÀN GIAO SỐ 2 – CHỜ KIỂM THỬ CHÉO
Bên giao tester: Bùi Xuân Minh –23110193. Bên nhận chủ hệ thống bạn cặp: '''+PEER+'''. Tên/link hệ thống bạn cặp: chưa được cung cấp. Trạng thái: CHƯA THỰC HIỆN, không ký nhận.
## I. Danh mục dự kiến
Phiếu review SRS,TestPlan,Testcase/Log/BugReport,Postman/Newman,Puppeteer,JMeter/JTL/HTML,Summary. Hiện có bộ tự kiểm thử StudentCare để tham chiếu cách làm; không đưa số của hệ thống mình vào ô kết quả bạn cặp.
## II. Tóm tắt kết quả bạn cặp
Functional/NFR: chưa chạy; Passrate: chưa xác định; Bugtheo5loại: chưa xác định; tỉlệpháthiệnseed: chưa xác định do chưa nhậnManifesttừgiảngviên saukiểmthử; kết luận: chưa đủdữliệu,khôngGo.
## III. Bất đồng và chữ ký
Chưa ghi nhận ý kiến của bên nhận; không đồng nghĩa “không có bất đồng”. BÊN GIAO/BÊN NHẬN: chưa ký. Cần thông tin bàn giao thực trước khi hoàn tất biểu mẫu này.
'''
    write_document('Mau_06_Bien_ban_ban_giao_2_CHO_BAN_CAP_23110193','BIÊN BẢN KẾT QUẢ KIỂM THỬ – MẪU 06 (CHƯA THỰC HIỆN)',hand2,PRIVATE)
    score='## Thông tin\nSinh viên: '+IDENTITY+'. Hệ thống mình: StudentCare. Bạn cặp: '+PEER+'. Ngàybảovệ: chờ lịchgiảngviên.\n## Bảngchấm\n'+table(['Mã','Tiêu chí','Tối đa','Điểm','Nhận xét'],[(a,b,c,'Giảng viên điền','Chưa chấm') for a,b,c in [('A1','URD',1),('A2','SRS',1.5),('A3','API+sảnphẩm',.5),('A4','Lỗicàisẵn+manifest',1),('B1','Review+testplan',.5),('B2','Testcase',1.5),('B3','Bugreport+tỉlệpháthiện',1.5),('B4','Postman+Puppeteer',1),('B5','JMeter',1.5)]])+ '\nCộng tối đa10; điểm cộng/trừ,điều chỉnhsau bảo vệ và chữký: GIẢNG VIÊN ĐIỀN. Không tự ghi điểm hay mạo chữký.\n'
    write_document('Phieu_cham_diem_23110193','PHIẾU CHẤM ĐỒ ÁN CÁ NHÂN (CHỜ GIẢNG VIÊN)',score,PRIVATE)


def build_summary():
    functional=load_json(BASE/'evidence/functional/training/results.json',{'total':0,'counts':{},'cases':[]})
    nfr=load_json(BASE/'evidence/nonfunctional/results.json',{'total':0,'counts':{},'cases':[]})
    newman=load_json(BASE/'evidence/tooling/newman/summary.json',{})
    puppeteer=load_json(BASE/'evidence/tooling/puppeteer/summary.json',{})
    tooling=load_json(BASE/'evidence/tooling/summary.json',{})
    bugs=load_json(BASE/'data/bug_catalog.json',[])
    if isinstance(bugs,dict):bugs=bugs.get('bugs',bugs.get('items',[]))
    totals=[]
    for key in ['total','PASS','FAIL','ERROR','NOT_RUN']:
        fv=functional['total'] if key=='total' else functional.get('counts',{}).get(key,0)
        nv=nfr['total'] if key=='total' else nfr.get('counts',{}).get(key,0)
        totals.append((key,fv,nv,fv+nv))
    body='''## Trạng thái của báo cáo
TỰ KIỂM THỬ StudentCare; tester và chủ hệ thống đều là '''+IDENTITY+'''. Không phải báo cáo độc lập của bạn cặp. Bạn cặp: '''+PEER+'''. Ngày thực thi lấy từ rawISOtimestamps; đợt bắt đầu24/10 và kết thúc25/10/2026 theoUTC+7. SRS1.1,TestPlan1.0; bản minh chứng có hashfile tại từng kết quả.
## 1. Tóm tắt
Đã kiểm tra16chức năng,3luồng và10NFR trên bản sao cục bộ. Các testcase dùng assertion hành vi đúng; FAIL được giữ khi phát hiện lỗi. Có lỗi chức năng/giao diện/nội dung/hiệu năng/HTTP500. Bản thực hành phù hợp để bạn cặp thử lại, nhưng chưa đủ điều kiện nghiệm thu nghiệp vụ. Phầnkiểmthửđộclập còn chờ thông tin bàn giao.
## 2. Kết quả thực thi
'''+table(['Chỉ số','Chức năng','Phi chức năng','Tổng'],totals)+'''
FAILtestcase không đồng nghĩa sốbug: nhiều E2E lặp lại cùngB06/B07/B10. Không dùng Passrate để tuyên bố hết lỗi. WorkbookTong_hop cócôngthức/trace và phầnbug phân loại riêng.
### 2.1. Độ phủ chức năng/NFR
'''+table(['Mã','Số TC','Pass','Fail','Chưa chạy'],[(fid,len([r for r in functional['cases'] if fid in r.get('requirements',[r.get('function_id')])]),sum(r.get('status')=='PASS' for r in functional['cases'] if fid in r.get('requirements',[r.get('function_id')])),sum(r.get('status')=='FAIL' for r in functional['cases'] if fid in r.get('requirements',[r.get('function_id')])),0) for fid,*_ in FR]+[(nid,len([r for r in nfr['cases'] if r['requirement']==nid]),sum(r['status']=='PASS' for r in nfr['cases'] if r['requirement']==nid),sum(r['status']=='FAIL' for r in nfr['cases'] if r['requirement']==nid),sum(r['status']=='NOT_RUN' for r in nfr['cases'] if r['requirement']==nid)) for nid,*_ in NFR])+'''
## 3. Thống kê lỗi
Manifestriêng20seed:10Functional,3Visual,3Content,2Performance,2Crash; thêm3exploratoryngoàiPDFgốc,X01–X03. Severity dựa ảnhhưởng, không dùng lỗi0.00xperformanceđểgánCritical. Danh mụcchi tiết/bằngchứng ởprivateworkbook; không gửi kèm tài liệu bạn cặp. Đây là kiểm chứngchủhệthống, tỉlệpháthiệnđộclập chưa xác định.
## 4. Kiểm thử tự động
### 4.1. Postman/Newman
'''+table(['Chỉ số','Thực đo'],[(k,newman.get(k,'Chưa có')) for k in ['requests','iterations','assertions','assertionsPassed','assertionsFailed']])+'''
Lệnh tái hiện: python tools/run_submission_tooling.py (npmdeps/JMeterđãchuẩnbị); lệnhAPIriêng node submission/tooling/run-newman.cjs theoREADME công cụ. Collection,environment,CSVdata và reportHTML nằm tooling/newman. Phủđăngnhập,phiên,danhmục,tạo,danh sách,chitiết,hủy,thôngbáo,logout,kiểmthửâm; ≥6chứcnăng.
### 4.2. Puppeteer
Rawsummary sau đây được trích từ lần chạy thực tế; sốFAILkhôngbịẩn bởi orchestrator exit0. Dữliệu táchui-data.json,ảnhtrongevidence/tooling/puppeteer.
'''+json.dumps(puppeteer,ensure_ascii=False,indent=2)+'''
## 5. JMeter
### 5.1. Môi trườngđo
MáyWindowscùngserverWerkzeug,khôngphảiVPS; DBtạm10hồsơ trước đo;Python3.11.9,JDK21,JMeter5.6.3; tảiHTTPloopback. Không đăngkýkhunggiờtảihạtầngtrường; khônggửi stress tới sv23. Thời gianthật trong JTL vàsummary.
### 5.2. Kịchbản
Tảithường2users,ramp10s,300s. Tăngtải1→2→4,mỗibậc120s,ramp5s; timer180sample/phút nhóm, tốiđa mục tiêu3request/s. Mỗithreadloginmộtlần,cookieđộc lập; đoSearch/List/Notifications/Types,logintáchriêng.
### 5.3–5.4. Số đo và đối chiếuSRS
'''
    for group in ('jmeter-load','jmeter-step'):
        report=load_json(BASE/'evidence/tooling'/group/'summary.json')
        body+='### '+group+'\n'
        if report:
            body+=json.dumps(report,ensure_ascii=False,indent=2)+'\n'
        else:body+='CHƯA CÓ KẾT QUẢ HOÀN TẤT – không ghi Pass hay số đo dự đoán.\n'
    body+='''### 5.5. Nhận xét
Đọc ResponseTimesOverTime/Throughput/ActiveThreads trongdashboardgốc. SoP95Searchvới800ms,Notifications900ms,List1000ms vàerror<1%; không tính thinktimevàoHTTPelapsed. ĐộtrễSearch/Notificationscốýsẽkhôngđạt ngưỡng; dùng sốJTLthực thay vì suy từsleep. Chỉ đo đến4user; chưa đủcăn cứkếtluận điểm bãohòa hoặc nănglựcVPS. Giảthuyếtbottleneckđượcphânbiệtvớithựcđo.
## 6. Rủi ro/hạn chế
Chưa cókiểmthửđộclập bạncặp; chưa chạytải trênVPS; chưa kiểmđa trìnhduyệt;chưa rebootVPS; chưa biết lịchdemo. Manifest10lỗiphiênbảncũđãcôngkhai tronggit history,khôngthểcamkết blindtest retroactive; bảnmanifestmớiphảigửiriêng. Không tựgiảm sốlỗi để tăngpassrate. Nhữngphầnblockedghiđúngchưathựchiện.
## 7. Kếtluận
NO-GO nghiệmthu nghiệpvụ: cònMajor và cácNFRkhôngđạt. READY vềbộhồsơkỹthuậttựkiểmtra đểbạncặptiếpnhận; phầntesterchínhthứcCHƯAHOÀNTẤT nếu chưa có dữliệubạncặp/chữký. Không sửa lỗi cài sẵn trướcdemo; sự cốmôi trường được xửlýriêng.
## 8. Khai báo AI
Codex hỗtrợmãnguồn,chuẩn hóatàiliệu,thiếtkếcase,runner vàtổnghợpbằngchứng. API/UI/JMeterthựcsựđượcchạy,rawfileđínhkèm. SinhviênBùiXuânMinhcầnđọclại,chạyđốichiếu,thựchiệnkiểmthửchéovàthảoluận;khôngthaythếxácnhậncủabạncặp/giảngviên.
'''
    write_document('Mau_05_Test_Summary_TU_KIEM_THU_23110193','BÁO CÁO TỔNG KẾT – MẪU 05 (TỰ KIỂM THỬ)',body,SELF)


def main():
    build_urd_srs();build_api()
    build_plan_and_handover();build_summary()
    print('Generated public URD/SRS/API/handover and private self-test plan/summary/unsigned forms.')


if __name__=='__main__':main()
